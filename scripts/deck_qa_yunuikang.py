#!/usr/bin/env python3
"""덱 레이아웃 QA — 이 머신에 렌더러(LibreOffice)가 없어 기하로 대신 검사한다.

선언된 박스 크기가 아니라 **렌더링될 텍스트의 실제 범위를 추정**해서
(한글 1.0em / 라틴 0.52em, 박스 폭 기준 자동 줄바꿈 반영) 겹침·화면 넘침을 잡는다.
표는 python-pptx가 최소 행높이만 기록하므로 폰트·줄수로 실제 높이를 다시 계산한다.

사용법:
    python3 scripts/deck_qa_yunuikang.py slides/<파일>.pptx

"실질 문제: 0" 이 나올 때까지 고친다.  규격은 docs/DECK_STYLE_yunuikang.md §8.
"""
import math
import sys

from pptx import Presentation
from pptx.util import Emu

SW, SH = 13.333, 7.5          # 16:9 캔버스 (in)
TOL = 0.06                    # 이보다 작은 겹침은 무시 (추정 오차)
EDGE = 0.05                   # 화면 밖 판정 여유


def _text_extent(sh):
    """(폭, 높이) 인치 추정. 한글은 1.0em, 라틴은 0.52em로 잡는다."""
    box_w = Emu(sh.width).inches - 0.1
    total_h, max_w = 0.0, 0.0
    for p in sh.text_frame.paragraphs:
        txt = "".join(r.text for r in p.runs)
        size = max([r.font.size.pt for r in p.runs if r.font.size] or [13])
        em = size / 72.0
        w = sum((1.0 if ord(ch) > 0x1100 else 0.52) * em for ch in txt)
        lines = max(1, math.ceil(w / box_w)) if box_w > 0 else 1
        total_h += lines * em * 1.22
        max_w = max(max_w, min(w, box_w) if box_w > 0 else w)
    return max_w, total_h


def _table_extent(sh):
    """표는 선언 높이가 최소값일 뿐 — 폰트·줄수로 실제 높이를 다시 잡는다."""
    h = 0.0
    for row in sh.table.rows:
        nlines, fs = 1, 8.0
        for c in row.cells:
            nlines = max(nlines, len(c.text.split("\n")))
            for p in c.text_frame.paragraphs:
                for run in p.runs:
                    if run.font.size:
                        fs = max(fs, run.font.size.pt)
        h += nlines * (fs / 72.0) * 1.22 + 0.04
    return Emu(sh.width).inches, h


def check(path):
    prs = Presentation(path)
    bad = 0
    for i, s in enumerate(prs.slides, 1):
        items = []
        for sh in s.shapes:
            if sh.width is None or sh.top is None:
                continue
            L, T = Emu(sh.left).inches, Emu(sh.top).inches
            if sh.has_table:
                w, h = _table_extent(sh)
                items.append((L, T, w, h, "TABLE"))
            elif sh.has_text_frame and sh.text_frame.text.strip():
                w, h = _text_extent(sh)
                items.append((L, T, w, h, sh.text_frame.text.strip().replace("\n", "⏎")[:34]))
        for j, a in enumerate(items):
            if a[0] < -EDGE or a[1] < -EDGE or a[0] + a[2] > SW + EDGE or a[1] + a[3] > SH + EDGE:
                print(f"S{i} OVERFLOW {a[4]!r}  R={a[0]+a[2]:.2f} B={a[1]+a[3]:.2f}")
                bad += 1
            for b in items[j + 1:]:
                ox = min(a[0] + a[2], b[0] + b[2]) - max(a[0], b[0])
                oy = min(a[1] + a[3], b[1] + b[3]) - max(a[1], b[1])
                if ox > TOL and oy > TOL:
                    print(f'S{i} COLLIDE {a[4]!r} <-> {b[4]!r}  (겹침 {ox:.2f}×{oy:.2f}")')
                    bad += 1

    # 부수 점검: 리터럴 마크다운 · 노트 누락
    lit = 0
    for s in prs.slides:
        for sh in s.shapes:
            if sh.has_text_frame and "**" in sh.text_frame.text:
                lit += sh.text_frame.text.count("**")
            if sh.has_table:
                for row in sh.table.rows:
                    for c in row.cells:
                        lit += c.text.count("**")
    no_notes = [i for i, s in enumerate(prs.slides, 1)
                if not s.notes_slide.notes_text_frame.text.strip()]
    if lit:
        print(f"⚠️  리터럴 '**' {lit}개 — add_text/add_table 헬퍼를 우회했는지 확인")
        bad += 1
    if no_notes:
        print(f"⚠️  발표자 노트 없는 슬라이드: {no_notes}")
        bad += 1

    print(f"\n슬라이드 {len(prs.slides.__iter__.__self__._sldIdLst)}장 · 실질 문제: {bad}")
    return bad


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit("usage: deck_qa_yunuikang.py <deck.pptx>")
    sys.exit(1 if check(sys.argv[1]) else 0)
