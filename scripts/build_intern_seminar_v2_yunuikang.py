#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""인턴 세미나 발표 덱 v2 (10분) — 앞부분(흐름·배경·문제·해법)을 풀어서 재구성.
개정: ①로드맵 신설 ②Problem Definition→Background 축소 ③Related Work→문제(스래싱)+해법 2장
      ④locality/balancing 프레임 삭제 ⑤분석 도입부에 'compute 안 봄→연산 남아돎' 리드 추가.
모든 수치·코드라인은 logs/*.md 실측값만 사용(지어내기 금지).
Output: slides/seminar_260708_v2.pptx   ·   Font: Calibri(라틴)+Malgun Gothic(한글)
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn
from PIL import Image

FIGS = "/home/yunuikang/yunuikang_work/distserving/figures"
OUT  = "/home/yunuikang/yunuikang_work/distserving/slides/seminar_260708_v2.pptx"
os.makedirs(os.path.dirname(OUT), exist_ok=True)

INK   = RGBColor(0x1f, 0x24, 0x30)
BLUE  = RGBColor(0x27, 0x46, 0x90)
AMBER = RGBColor(0xE0, 0x8A, 0x1E)
GREEN = RGBColor(0x2E, 0x7D, 0x46)
RED   = RGBColor(0xC0, 0x39, 0x2B)
GRAY  = RGBColor(0x6B, 0x72, 0x80)
LGRAY = RGBColor(0xEC, 0xEE, 0xF2)
PANEL = RGBColor(0xF4, 0xF6, 0xFB)
BLUEL = RGBColor(0xE4, 0xEA, 0xF6)
AMBRL = RGBColor(0xFB, 0xF0, 0xDD)
GRNL  = RGBColor(0xE3, 0xF1, 0xE8)
REDL  = RGBColor(0xF8, 0xE4, 0xE1)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LATIN = "Calibri"
EA    = "Malgun Gothic"

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]

def slide():
    return prs.slides.add_slide(BLANK)

def _set_font(run, size, color, bold=False, italic=False):
    run.font.size = Pt(size); run.font.color.rgb = color
    run.font.bold = bold; run.font.italic = italic
    run.font.name = LATIN
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {}); rPr.append(el)
        el.set("typeface", EA if tag == "a:ea" else LATIN)

def textbox(s, l, t, w, h, lines, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, sa=3):
    tb = s.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = Pt(4); tf.margin_right = Pt(4); tf.margin_top = Pt(2); tf.margin_bottom = Pt(2)
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        for (txt, sz, col, bold) in (ln if isinstance(ln, list) else [ln]):
            r = p.add_run(); r.text = txt; _set_font(r, sz, col, bold)
        p.space_after = Pt(sa); p.space_before = Pt(0)
    return tb

def rect(s, l, t, w, h, fill, line=None, lw=1.0, shape=MSO_SHAPE.RECTANGLE, radius=None):
    sp = s.shapes.add_shape(shape, Inches(l), Inches(t), Inches(w), Inches(h))
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line; sp.line.width = Pt(lw)
    sp.shadow.inherit = False
    return sp

def boxtext(s, l, t, w, h, lines, fill, line=None, lw=1.0, align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    rect(s, l, t, w, h, fill, line, lw, shape=shape)
    textbox(s, l, t, w, h, lines, align=align, anchor=anchor)

def arrow(s, x1, y1, x2, y2, color=GRAY, lw=2.25):
    cn = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    cn.line.color.rgb = color; cn.line.width = Pt(lw)
    le = cn.line._get_or_add_ln()
    tail = le.makeelement(qn('a:tailEnd'), {'type': 'triangle', 'w': 'med', 'len': 'med'})
    le.append(tail)
    cn.shadow.inherit = False
    return cn

def title(s, tag, tag_color, head):
    rect(s, 0.55, 0.42, 0.14, 0.62, tag_color)
    textbox(s, 0.80, 0.34, 12.0, 0.66, [[(head, 25, INK, True)]], anchor=MSO_ANCHOR.MIDDLE)
    textbox(s, 0.80, 1.00, 12.0, 0.34, [[(tag, 12.5, tag_color, True)]])
    rect(s, 0.80, 1.40, 11.95, 0.02, tag_color)

def add_image(s, name, top, width=None, center=True, left=None, max_h=None):
    path = os.path.join(FIGS, name)
    iw, ih = Image.open(path).size
    if width is None: width = 11.7
    w = Inches(width); h = Emu(int(w * ih / iw))
    if max_h is not None and h > Inches(max_h):
        h = Inches(max_h); w = Emu(int(h * iw / ih))
    if left is None:
        left = (SW - w) / 2 if center else Inches(0.8)
    else:
        left = Inches(left)
    s.shapes.add_picture(path, left, Inches(top), width=w, height=h)
    return Inches(top) + h

def notes(s, text):
    s.notes_slide.notes_text_frame.text = text

def table(s, data, l, t, w, col_w, header_fill=BLUE, fs=12, hfs=12, row_h=0.42,
          zebra=True, cell_colors=None):
    rows, cols = len(data), len(data[0])
    gt = s.shapes.add_table(rows, cols, Inches(l), Inches(t), Inches(w), Inches(row_h*rows)).table
    gt.first_row = False; gt.horz_banding = False
    for ci, cw in enumerate(col_w):
        gt.columns[ci].width = Inches(cw)
    for ri in range(rows):
        gt.rows[ri].height = Inches(row_h)
        for ci in range(cols):
            cell = gt.cell(ri, ci)
            cell.margin_left = Pt(6); cell.margin_right = Pt(6)
            cell.margin_top = Pt(2); cell.margin_bottom = Pt(2)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            if ri == 0:
                cell.fill.solid(); cell.fill.fore_color.rgb = header_fill
            elif cell_colors and cell_colors.get((ri, ci)):
                cell.fill.solid(); cell.fill.fore_color.rgb = cell_colors[(ri, ci)]
            elif zebra and ri % 2 == 0:
                cell.fill.solid(); cell.fill.fore_color.rgb = PANEL
            else:
                cell.fill.solid(); cell.fill.fore_color.rgb = WHITE
            txt, col, bold = data[ri][ci]
            tf = cell.text_frame; tf.word_wrap = True
            p = tf.paragraphs[0]; p.alignment = PP_ALIGN.LEFT
            r = p.add_run(); r.text = txt
            _set_font(r, hfs if ri == 0 else fs, WHITE if ri == 0 else col, bold or ri == 0)
    return gt

def C(txt, col=INK, bold=False):
    return (txt, col, bold)

def footer(s, txt):
    textbox(s, 0.80, 7.06, 11.95, 0.34, [[(txt, 11, GRAY, False)]])

# ===========================================================================
# 1 — Title (유지)
# ===========================================================================
s = slide()
rect(s, 0, 0, 13.333, 7.5, WHITE)
rect(s, 0, 2.62, 13.333, 0.055, AMBER)
textbox(s, 0.9, 1.05, 11.6, 1.6,
        [[("이종(Heterogeneous) GPU 클러스터에서의", 30, BLUE, True)],
         [("에이전트 서빙 라우팅", 30, BLUE, True)]])
textbox(s, 0.9, 2.85, 11.6, 1.4,
        [[("SOTA 스케줄러 ", 15, GRAY, False), ("ThunderAgent", 15, AMBER, True),
          ("를 baseline으로, 이종 GPU·다양한 워크로드에서", 15, GRAY, False)],
         [("언제·왜 성능이 갈리는지를 하나의 인과 모델로 설명한다.", 15, GRAY, False)]])
textbox(s, 0.9, 6.25, 11.6, 0.9,
        [[("강윤의 (Yunui Kang)", 16, INK, True)],
         [("인턴 세미나  ·  개별연구 중간 발표  ·  2026-07-08", 12.5, GRAY, False)]])
notes(s,
"안녕하세요. 개별연구를 진행하고 있는 강윤의입니다. 오늘은 10분 동안 제 개별연구를 소개하겠습니다. "
"한 줄로 말씀드리면, 제 주제는 '서로 다른 종류의 GPU가 섞여 있는 클러스터에서, 들어오는 에이전트 요청을 "
"어느 GPU로 보낼지'를 연구하는 것입니다. 발표는 네 부분입니다. 먼저 제가 지금 전체 흐름에서 어디쯤 와 "
"있는지 로드맵을 보여드리고, 그 다음 배경과 문제, 그리고 baseline인 ThunderAgent가 무엇을 어떻게 푸는지, "
"마지막으로 실험 결과와 원인, 향후 계획을 말씀드리겠습니다. 핵심 메시지 하나만 미리 — 흩어져 있던 실험 "
"결과들을 'R'이라는 하나의 값으로 언제 이기고 언제 지는지 예측할 수 있게 됐다는 것입니다.")
footer(s, "개별연구 · 이종 GPU 에이전트 서빙")

# ===========================================================================
# 2 — 로드맵 (신설) ★최우선
# ===========================================================================
s = slide()
title(s, "현재 진행 상황 · Where We Are", BLUE, "지금 우리는 'baseline으로 병목을 찾는' 단계에 있다")
steps = [
    ("1", "방향 설정", "Distributed 환경\n에이전트 서빙", False),
    ("2", "관련 연구 조사", "SOTA =\nThunderAgent", False),
    ("3", "병목 탐색", "Baseline으로 돌려\nbottleneck 찾기", True),
    ("4", "환경 변경 실험", "이종 GPU 및\n다양한 워크로드", False),
    ("5", "분석 · 향후", "결과·원인 분석\n→ 향후 방향", False),
]
bx, bw, gap = 0.72, 2.15, 0.36
ytop, bh = 2.85, 1.85
for i, (n, tag, body, cur) in enumerate(steps):
    x = bx + i*(bw+gap)
    if cur:
        textbox(s, x-0.25, ytop-0.52, bw+0.5, 0.42, [[("▼ 지금 여기", 13, AMBER, True)]], align=PP_ALIGN.CENTER)
    rect(s, x, ytop, bw, bh, AMBRL if cur else PANEL, line=AMBER if cur else LGRAY,
         lw=2.4 if cur else 1.2, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    rect(s, x, ytop, bw, 0.44, AMBER if cur else BLUE)
    textbox(s, x, ytop, bw, 0.44, [[(f"{n}. {tag}", 12, WHITE, True)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    textbox(s, x+0.10, ytop+0.52, bw-0.20, 1.25,
            [[(ln, 12, INK, cur)] for ln in body.split("\n")], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    if i < 4:
        arrow(s, x+bw, ytop+bh/2, x+bw+gap, ytop+bh/2, GRAY, 2.5)
rect(s, 0.90, 5.35, 11.85, 1.35, PANEL, line=BLUE, lw=1.3, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
textbox(s, 1.15, 5.48, 11.4, 1.15,
        [[("큰 그림 — ", 14, BLUE, True),
          ("서로 다른 GPU가 섞인 클러스터에서 에이전트 요청을 어디로 보낼지를 연구한다.", 14, INK, False)],
         [("맨땅에서 새로 만들지 않는다: 이 분야 SOTA인 ThunderAgent를 baseline으로 돌려 병목을 먼저 찾고,",
           12.5, INK, False)],
         [("환경(이종 GPU·워크로드)을 바꿔 실험한 뒤 원인을 규명하는 흐름 — 오늘은 3~5단계를 보고드린다.",
           12.5, INK, False)]])
notes(s,
"본론에 들어가기 전에, 제가 전체 흐름에서 지금 어디쯤 있는지부터 보여드리겠습니다. 저는 distributed 서빙"
"이라는 큰 방향을 잡고 시작했습니다. 1단계로 방향을 정하고, 2단계로 관련 최신 연구를 조사해서 이 분야의 "
"SOTA인 ThunderAgent를 찾았습니다. 지금은 3단계, 그 ThunderAgent를 baseline으로 실제로 돌려보면서 어디서 "
"막히는지 병목을 찾는 단계에 있습니다. 그래서 4단계로 환경을 바꿔 — 이종 GPU와 여러 워크로드로 — 실험했고, "
"5단계로 그 결과와 원인을 분석하고 있습니다. 오늘 발표는 이 3단계부터 5단계까지, 즉 baseline을 돌려 병목을 "
"찾고 원인을 설명한 내용이 중심입니다. 핵심은, 맨땅에서 시작한 게 아니라 SOTA를 딛고 그 한계를 찾는 "
"과정이라는 점입니다.")
footer(s, "① 방향 → ② SOTA 조사 → ③ baseline 병목 탐색(현재) → ④ 환경 변경 실험 → ⑤ 분석·향후")

# ===========================================================================
# 3 — Background (Problem Definition 축소): 에이전트 서빙의 특성
# ===========================================================================
s = slide()
title(s, "Background", GRAY, "에이전트 서빙의 특성 — 왜 '요청'이 아니라 '프로그램' 단위인가")
textbox(s, 0.80, 1.55, 11.95, 0.5,
        [[("한 문장: ", 14.5, INK, True),
          ("에이전트는 리즈닝↔툴콜을 여러 턴 반복 → 턴마다 KV 캐시가 쌓인다 → 요청 하나가 아니라 프로그램 전체를 관리해야 한다.",
           14.5, INK, False)]])
# ReAct 루프 (왼쪽)
lx, ly = 1.05, 2.55
boxtext(s, lx, ly, 2.15, 0.78, [[("REASONING", 13, WHITE, True)], [("GPU 추론", 10.5, BLUEL, False)]], BLUE)
boxtext(s, lx+3.15, ly, 2.15, 0.78, [[("ACTING", 13, WHITE, True)], [("툴 콜 (GPU idle)", 10.5, AMBRL, False)]], AMBER)
arrow(s, lx+2.15, ly+0.39, lx+3.15, ly+0.39, GRAY, 2.2)
arrow(s, lx+3.15, ly+0.66, lx+2.15, ly+0.66, GRAY, 2.2)
textbox(s, lx-0.15, ly+0.92, 5.7, 0.36, [[("1턴 = 리즈닝+툴콜 · 멀티턴 반복", 11, GRAY, False)]], align=PP_ALIGN.CENTER)
# KV 누적 막대 (왼쪽 하단)
textbox(s, lx-0.15, ly+1.45, 5.7, 0.30, [[("턴이 쌓일수록 KV 캐시 누적 →", 12, INK, True)]])
for i in range(4):
    hh = 0.24 + i*0.20
    rect(s, lx + i*1.25, ly+2.85-hh, 0.95, hh, BLUEL, line=BLUE, lw=1.0)
    textbox(s, lx + i*1.25, ly+2.88, 0.95, 0.3, [[(f"T{i+1}", 10, GRAY, False)]], align=PP_ALIGN.CENTER)
# 결론 카드 (오른쪽)
rect(s, 7.15, 2.55, 5.6, 2.95, GRNL, line=GREEN, lw=1.6, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
textbox(s, 7.4, 2.70, 5.15, 2.7,
        [[("챗봇형 LLM", 13, GRAY, True), ("  :  입력 1회 → 출력 1회 (stateless)", 12, INK, False)],
         [("에이전트", 13, GREEN, True), ("  :  멀티턴, KV가 프로그램 생애 내내 누적 (stateful)", 12, INK, False)],
         [("", 8, INK, False)],
         [("→ 핵심 결론", 14, GREEN, True)],
         [("스케줄링 단위가 ", 13, INK, False), ("요청(request)", 13, RED, True),
          ("이 아니라", 13, INK, False)],
         [("프로그램(program)", 15, GREEN, True), (" 이어야 한다.", 13, INK, True)],
         [("(이 캐시를 어느 GPU에 두고 언제 뺄지가 성능을 지배)", 11, GRAY, False)]])
notes(s,
"본격적인 배경입니다. 기존 챗봇형 LLM은 입력 한 번에 출력 한 번이면 끝나는, 상태가 없는 작업입니다. "
"그런데 에이전트는 다릅니다. 왼쪽 그림처럼 리즈닝과 툴 콜을 왔다갔다 반복합니다. 리즈닝은 GPU를 쓰지만 "
"툴 콜은 외부 API나 코드 실행이라 그동안 GPU가 놉니다. 그리고 중요한 건, 이게 여러 턴 반복되면서 왼쪽 "
"아래처럼 KV 캐시가 턴마다 계속 쌓인다는 점입니다. 앞 대화가 전부 컨텍스트로 붙기 때문입니다. 그래서 "
"오른쪽 결론이 나옵니다 — 챗봇은 요청 하나가 독립적이지만, 에이전트는 한 프로그램이 여러 턴에 걸쳐 상태를 "
"쌓아갑니다. 따라서 스케줄링을 요청 단위가 아니라 프로그램 단위로 해야 합니다. 이 캐시를 어느 GPU에 두고 "
"언제 뺄지가 성능을 좌우하고, 그게 바로 다음 장에서 볼 ThunderAgent가 푸는 문제입니다.")
footer(s, "리즈닝↔툴콜 멀티턴 → KV가 프로그램 생애 내내 누적 → 요청이 아닌 '프로그램' 단위 관리 필요")

# ===========================================================================
# 4 — ThunderAgent가 푸는 문제 = KV 스래싱 (신설, from Related Work)
# ===========================================================================
s = slide()
title(s, "ThunderAgent가 푸는 문제", RED, "KV 캐시 스래싱 — 동시성이 오르면 서로의 캐시를 지워 재프리필이 폭증한다")
# 시나리오 (상단, native 4-box)
sy = 1.65
scen = [
    ("P1 툴 콜 중", "GPU 잠깐 양보", AMBER),
    ("P2가 GPU 필요", "→ P1의 KV를 evict", RED),
    ("P1 다음 턴", "캐시 없음 → 전체 재프리필", RED),
    ("반복 = 스래싱", "재프리필 폭증", INK),
]
sbw, sgap = 2.72, 0.42
sxx = 0.80
for i, (h, d, col) in enumerate(scen):
    x = sxx + i*(sbw+sgap)
    rect(s, x, sy, sbw, 1.0, PANEL, line=col, lw=1.6, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    textbox(s, x+0.12, sy+0.10, sbw-0.24, 0.44, [[(h, 12.5, col, True)]], align=PP_ALIGN.CENTER)
    textbox(s, x+0.12, sy+0.52, sbw-0.24, 0.42, [[(d, 10.5, INK, False)]], align=PP_ALIGN.CENTER)
    if i < 3:
        arrow(s, x+sbw, sy+0.5, x+sbw+sgap, sy+0.5, GRAY, 2.2)
# 실측 그래프 (하단)
add_image(s, "fig_thrash_default.png", 3.15, width=8.7)
textbox(s, 0.80, 6.55, 11.95, 0.4,
        [[("실측(2×4090): 동시성 8→48에서 p95 지연 5.6× 폭증, 캐시 적중률 붕괴 → throughput은 천장에 막힘.",
           11.5, GRAY, True)]], align=PP_ALIGN.CENTER)
notes(s,
"이제 baseline인 ThunderAgent가 정확히 무슨 문제를 푸는지입니다. 그 문제가 바로 KV 캐시 스래싱입니다. "
"위쪽 시나리오를 보시면 — 프로그램 P1이 툴 콜을 하느라 GPU를 잠깐 비웁니다. 그 사이 다른 프로그램 P2가 "
"GPU가 필요해서 자리를 확보하려고 P1의 KV 캐시를 지워버립니다. 그런데 P1의 다음 턴이 오면, 캐시가 없으니 "
"그 긴 대화를 처음부터 전부 다시 계산해야 합니다. 이게 여러 프로그램 사이에서 반복되면 재프리필이 폭증하는 "
"스래싱이 됩니다. 아래는 이걸 실제로 측정한 그래프입니다. 요청 단위로만 처리하는 naive 방식에서 동시성을 "
"8에서 48로 올리면, 왼쪽처럼 p95 지연이 5.6배로 폭증하고, 오른쪽처럼 캐시 적중률이 0으로 붕괴합니다. "
"처리량은 더 올리려 해도 천장에 막힙니다. 즉 동시에 돌리는 프로그램이 많아질수록 성능이 무너지는 게 "
"핵심 병목이고, ThunderAgent는 바로 이걸 푸는 연구입니다.")
footer(s, "스래싱: 툴콜 중 evict → 다음 턴 전체 재프리필 → 동시성↑일수록 지연 폭증·throughput 천장")

# ===========================================================================
# 5 — ThunderAgent의 해법 = 프로그램 단위 스케줄링 (신설)
# ===========================================================================
s = slide()
title(s, "ThunderAgent의 해법", GREEN, "프로그램 단위 스케줄링 — 툴콜 중인 프로그램을 pause/resume해 evict를 사전 차단")
# Row 1 (기존 request 단위, 나쁨)
r1y = 1.75
rect(s, 0.80, r1y, 11.95, 1.35, REDL, line=RED, lw=1.4, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
textbox(s, 0.98, r1y+0.10, 3.0, 1.15, [[("기존: 요청 단위", 13, RED, True)], [("(엔진이 프로그램을", 10.5, GRAY, False)], [(" 모름)", 10.5, GRAY, False)]], anchor=MSO_ANCHOR.MIDDLE)
r1 = [("자리 부족", RED), ("남의 KV evict", RED), ("다음 턴 recompute", RED), ("스래싱", INK)]
for i, (t, col) in enumerate(r1):
    x = 4.05 + i*2.05
    rect(s, x, r1y+0.36, 1.72, 0.62, WHITE, line=col, lw=1.2, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    textbox(s, x, r1y+0.36, 1.72, 0.62, [[(t, 11, col, True)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    if i < 3: arrow(s, x+1.72, r1y+0.67, x+2.05, r1y+0.67, GRAY, 1.8)
# Row 2 (program 단위, 좋음)
r2y = 3.35
rect(s, 0.80, r2y, 11.95, 1.35, GRNL, line=GREEN, lw=1.6, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
textbox(s, 0.98, r2y+0.10, 3.0, 1.15, [[("해법: 프로그램 단위", 13, GREEN, True)], [("(ThunderAgent)", 10.5, GRAY, False)]], anchor=MSO_ANCHOR.MIDDLE)
r2 = [("용량 초과 감지", GREEN), ("작은 프로그램 pause", AMBER), ("큐에서 캐시 보존", GREEN), ("여유 시 resume", GREEN)]
for i, (t, col) in enumerate(r2):
    x = 4.05 + i*2.05
    rect(s, x, r2y+0.36, 1.72, 0.62, WHITE, line=col, lw=1.2, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    textbox(s, x, r2y+0.36, 1.72, 0.62, [[(t, 10.5, col, True)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    if i < 3: arrow(s, x+1.72, r2y+0.67, x+2.05, r2y+0.67, GRAY, 1.8)
textbox(s, 4.05, r2y+1.02, 8.5, 0.3, [[("→ evict·recompute 자체가 안 일어남 (캐시 보존)", 11, GREEN, True)]])
# 하단: homo→hetero 확장
rect(s, 0.80, 5.15, 11.95, 1.4, PANEL, line=AMBER, lw=1.4, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
textbox(s, 1.05, 5.28, 11.5, 1.2,
        [[("남은 한계 → 우리 연구의 출발점", 13.5, AMBER, True)],
         [("ThunderAgent를 포함한 기존 연구는 대부분 ", 12.5, INK, False),
          ("동종(homogeneous) GPU", 12.5, RED, True), ("를 가정한다.", 12.5, INK, False)],
         [("→ 우리는 ", 12.5, INK, False), ("이종(4090+5090) GPU + 다양한 워크로드", 12.5, GREEN, True),
          ("로 넣었을 때 어떤 병목이 드러나는지를 규명한다 (다음 장부터).", 12.5, INK, False)]])
notes(s,
"그럼 ThunderAgent는 이 스래싱을 어떻게 풀까요? 핵심 아이디어 하나입니다 — 요청이 아니라 프로그램 단위로 "
"스케줄링하는 것입니다. 위 빨간 줄이 기존 방식입니다. 자리가 부족하면 남의 캐시를 지우고, 그 프로그램의 "
"다음 턴에 다시 계산하고, 그게 스래싱이 됩니다. 아래 초록 줄이 ThunderAgent입니다. 메모리 용량이 초과될 "
"것 같으면, 지금 툴 콜 중이라 GPU를 안 쓰는 작은 프로그램을 잠깐 빼둡니다. 이걸 pause라고 합니다. 그 "
"프로그램의 캐시는 지우는 게 아니라 큐에 보존해 두고, 자리가 나면 다시 넣습니다. 이게 resume입니다. 그래서 "
"evict와 재계산 자체가 일어나지 않습니다. 세부 알고리즘은 뒤 메커니즘 장에서 보고, 여기선 직관만 짚습니다. "
"마지막으로 아래 주황 박스가 제 연구의 출발점입니다. ThunderAgent를 포함한 기존 연구는 대부분 모든 GPU가 "
"똑같다고 가정합니다. 하지만 현실은 4090, 5090처럼 섞여 있죠. 그래서 저는 이종 GPU와 다양한 워크로드에 "
"넣었을 때 어떤 병목이 새로 드러나는지를 봤고, 그게 다음 장부터입니다.")
footer(s, "해법 = 프로그램 단위 pause/resume로 evict 사전 차단 · 단 기존은 동종 가정 → 이종 확장이 우리 출발점")

# ===========================================================================
# 6 — System Overview (단순화: locality/balancing 프레임 삭제)
# ===========================================================================
s = slide()
title(s, "System Overview", BLUE, "Client → Router → 이종 인스턴스 (4090 / 5090)")
# Client
boxtext(s, 0.75, 2.55, 1.7, 1.0, [[("Client", 14, WHITE, True)], [("에이전트\n프로그램", 10, BLUEL, False)]],
        INK, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
arrow(s, 2.45, 3.05, 3.35, 3.05, GRAY, 2.5)
# Router
rect(s, 3.35, 2.20, 3.55, 2.30, PANEL, line=BLUE, lw=2.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
textbox(s, 3.35, 2.30, 3.55, 0.5, [[("Router / Orchestrator", 13.5, BLUE, True)]], align=PP_ALIGN.CENTER)
for i, (t, c, b) in enumerate([("① 신규 프로그램 라우팅", INK, True), ("② pause / resume", AMBER, True),
                               ("③ 용량·부하 판단", INK, True)]):
    textbox(s, 3.5, 2.86+i*0.46, 3.3, 0.4, [[(t, 12, c, b)]], align=PP_ALIGN.LEFT)
# arrows to instances
arrow(s, 6.9, 2.85, 8.15, 2.40, BLUE, 2.5)
arrow(s, 6.9, 3.75, 8.15, 4.40, BLUE, 2.5)
arrow(s, 8.15, 3.15, 6.9, 3.15, AMBER, 1.75)
textbox(s, 6.95, 2.80, 1.25, 0.3, [[("route", 9.5, GRAY, False)]], align=PP_ALIGN.CENTER)
textbox(s, 6.95, 3.22, 1.25, 0.3, [[("pause↩", 9.5, AMBER, True)]], align=PP_ALIGN.CENTER)
# Instance 4090
rect(s, 8.20, 1.80, 4.55, 1.30, WHITE, line=GRAY, lw=1.5, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
textbox(s, 8.35, 1.88, 4.3, 0.4, [[("Instance A — RTX 4090", 13, INK, True)]])
rect(s, 8.35, 2.30, 1.55, 0.62, BLUEL, line=BLUE, lw=1.0)
textbox(s, 8.35, 2.33, 1.55, 0.56, [[("KV pool", 9.5, BLUE, True)], [("43,888 tok", 10.5, INK, True)]], align=PP_ALIGN.CENTER)
textbox(s, 10.05, 2.29, 2.6, 0.7, [[("작은 용량 · 낮은 대역폭", 10.5, GRAY, False)],
                                     [("resident ≈ 2 프로그램", 10.5, GRAY, False)]])
# Instance 5090
rect(s, 8.20, 3.55, 4.55, 1.55, WHITE, line=GREEN, lw=1.8, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
textbox(s, 8.35, 3.63, 4.3, 0.4, [[("Instance B — RTX 5090", 13, INK, True)]])
rect(s, 8.35, 4.05, 2.05, 0.72, GRNL, line=GREEN, lw=1.0)
textbox(s, 8.35, 4.10, 2.05, 0.64, [[("KV pool", 9.5, GREEN, True)], [("89,040 tok", 11, INK, True)]], align=PP_ALIGN.CENTER)
textbox(s, 10.55, 4.02, 2.15, 0.9, [[("큰 용량(2.03×)", 10.5, GREEN, True)],
                                      [("높은 처리량/대역폭", 10.5, GRAY, False)],
                                      [("resident ≈ 4", 10.5, GRAY, False)]])
# 상태기계 (하단 좌)
rect(s, 0.90, 5.35, 6.0, 1.55, PANEL, line=LGRAY, lw=1.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
textbox(s, 1.05, 5.42, 5.7, 0.35, [[("요청 생애주기 상태기계", 12, INK, True)]])
boxtext(s, 1.05, 5.85, 1.55, 0.62, [[("ACTIVE", 11.5, WHITE, True)]], GREEN)
boxtext(s, 3.05, 5.85, 1.55, 0.62, [[("PAUSED", 11.5, WHITE, True)]], AMBER)
boxtext(s, 5.05, 5.85, 1.55, 0.62, [[("resume", 11.5, WHITE, True)]], BLUE)
arrow(s, 2.60, 6.16, 3.05, 6.16, GRAY, 2.0)
arrow(s, 4.60, 6.16, 5.05, 6.16, GRAY, 2.0)
textbox(s, 1.05, 6.50, 5.7, 0.35, [[("Eq.6(용량) 위반 시 pause → 프록시 큐 대기 → 여유 생기면 resume", 10, GRAY, False)]])
# 핵심 콜아웃 (하단 우) — locality/balancing 대신 이종 용량비
rect(s, 7.10, 5.35, 5.65, 1.55, GRNL, line=GREEN, lw=1.3, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
textbox(s, 7.25, 5.45, 5.4, 1.4,
        [[("이 슬라이드의 핵심 한 가지", 12, GREEN, True)],
         [("두 GPU의 KV 풀이 다르다 — ", 11, INK, False), ("5090이 4090의 2.03×", 11, GREEN, True)],
         [("→ '큰 GPU를 어떻게 채우느냐'가 뒤 분석·향후 과제의 열쇠", 11, INK, True)]])
notes(s,
"교수님이 그림을 좋아하셔서 시스템 구조를 한 장에 정리했습니다. 왼쪽부터, 클라이언트가 에이전트 프로그램을 "
"보내면 가운데 Router가 받습니다. Router가 하는 결정은 세 가지입니다 — 새 프로그램을 어느 GPU로 보낼지, "
"메모리가 부족하면 어떤 프로그램을 잠시 빼고 나중에 다시 넣을지, 그리고 각 GPU의 용량과 부하를 어떻게 "
"볼지입니다. 오른쪽이 인스턴스인데, 여기가 제 연구의 핵심입니다 — 4090은 KV 풀이 약 4만4천 토큰이라 "
"프로그램 두 개 정도만 올라가고, 5090은 약 8만9천 토큰으로 2배 크고 처리량도 높습니다. 아래 왼쪽은 요청의 "
"상태 변화입니다. 용량 조건을 위반하면 프로그램이 PAUSED로 큐에서 대기하다가 자리가 나면 resume됩니다. "
"아래 오른쪽이 오늘 꼭 기억하실 한 가지입니다 — 두 GPU의 KV 풀 크기가 2.03배 다르고, 이 큰 GPU를 어떻게 "
"채우느냐가 뒤에서 볼 분석과 향후 과제의 핵심 열쇠입니다.")
footer(s, "이종 인스턴스: 4090(43,888 tok) vs 5090(89,040 tok, 2.03×) · Router가 라우팅·pause/resume·용량 판단")

# ===========================================================================
# 7 — 실험 세팅 (유지)
# ===========================================================================
s = slide()
title(s, "실험 세팅", BLUE, "동종 2×4090에서 워크로드(duty)만 바꿔가며 tr vs default")
boxtext(s, 0.90, 1.75, 3.75, 1.65,
        [[("무엇을 비교?", 13, BLUE, True)],
         [("tr", 13, GREEN, True), ("  = ThunderAgent 스케줄러", 11.5, INK, False)],
         [("     (프로그램 단위 pause/resume)", 10, GRAY, False)],
         [("default", 13, RED, True), (" = 최소부하로 프록시만", 11.5, INK, False)],
         [("     (프로그램 통제 없음)", 10, GRAY, False)]],
        PANEL, line=BLUE, lw=1.3, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP)
boxtext(s, 4.85, 1.75, 3.75, 1.65,
        [[("하드웨어 (Hardware)", 13, BLUE, True)],
         [("2 × RTX 4090 (동종)", 11.5, INK, True)],
         [("KV 풀 43,888 tok / GPU", 10.5, GRAY, False)],
         [("이종(4090+5090) 용량차이는", 10.5, GRAY, False)],
         [("   메커니즘 분석·향후 과제로", 10.5, GRAY, False)]],
        PANEL, line=BLUE, lw=1.3, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP)
boxtext(s, 8.80, 1.75, 3.95, 1.65,
        [[("워크로드 — duty(d)로 대비", 13, BLUE, True)],
         [("합성 — balanced  d≈0.56", 11, INK, False)],
         [("TraceLab — prefill-heavy  d≈0.18", 11, INK, False)],
         [("SWE-bench — decode-heavy  d≈1.0", 11, INK, True)],
         [("d = GPU 실작동 시간 비율", 10, GRAY, False)]],
        PANEL, line=BLUE, lw=1.3, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP)
data = [
 [C("실험", WHITE, True), C("워크로드 (성격, duty d)", WHITE, True), C("HW", WHITE, True), C("이 실험이 답하는 것", WHITE, True)],
 [C("합성", GREEN, True), C("balanced (d≈0.56)"), C("2×4090"), C("tr이 의도대로 작동하는가 (기대)")],
 [C("D", BLUE, True), C("TraceLab · prefill-heavy (d≈0.18)"), C("2×4090"), C("실데이터에서도 그런가")],
 [C("C", GREEN, True), C("TraceLab (duty 스윕)"), C("격리 2×4090"), C("★ 승패 원인의 인과 증명 (R 모델)")],
 [C("SWE", GREEN, True), C("SWE-bench · decode-heavy (d≈1.0)"), C("2×4090"), C("★ 2차 실데이터로 결론 재현")],
]
table(s, data, 0.90, 3.80, 11.85, [1.0, 4.9, 1.3, 4.65], fs=12, hfs=12, row_h=0.54)
textbox(s, 0.90, 6.60, 11.85, 0.34,
        [[("모두 동종 2×4090 — 워크로드 duty만 바꾼다. 이종(5090) 용량차이 문제는 메커니즘(코드)·향후 과제에서 다룸.",
           10.5, GRAY, True)]])
notes(s,
"실험 세팅입니다. 비교하는 두 가지는, tr — ThunderAgent 스케줄러로 프로그램 단위로 pause와 resume을 하는 "
"방식이고, default — 부하 적은 GPU로 요청만 넘기고 프로그램 통제는 안 하는 방식입니다. 오늘 실험은 하드웨어를 "
"동종 2×4090으로 고정하고, 워크로드만 바꿔가며 봤습니다. 이렇게 하면 승패를 가르는 변수를 하드웨어가 아니라 "
"워크로드 하나로 깔끔하게 분리할 수 있습니다. 워크로드는 duty, 즉 전체 시간 중 GPU가 실제로 도는 비율로 "
"구분합니다. 합성은 중간, TraceLab은 입력이 길어 프리필이 무거운데 duty는 낮고, SWE-bench는 코딩 에이전트라 "
"생성이 길어 duty가 거의 1입니다. 표를 보시면, 먼저 합성으로 tr이 기대대로 작동하는지 확인하고, TraceLab "
"실데이터로 실제로도 그런지 봅니다. 그다음 실험 C와 SWE가 오른쪽 별표, 오늘 강조할 인과 증명입니다. 이종 "
"4090+5090의 용량 차이 문제는 뒤 메커니즘 분석에서 코드로 짚고 향후 과제로 연결하겠습니다.")
footer(s, "동종 2×4090 고정 · 워크로드 duty(d)만 변화 · 합성→TraceLab→(원인)→SWE")

# ===========================================================================
# 8 — 결과 ① 합성 (유지)
# ===========================================================================
s = slide()
title(s, "결과 ①  합성 워크로드 · 2×4090", GREEN, "우리 기대대로 — tr이 스래싱을 억제해 throughput·hit 모두 이긴다")
add_image(s, "thrash_throughput.png", 1.70, width=5.7, left=0.75)
add_image(s, "thrash_hit_rate.png", 1.70, width=5.7, left=6.95)
rect(s, 0.90, 6.05, 5.6, 1.0, GRNL, line=GREEN, lw=1.2, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
textbox(s, 1.05, 6.12, 5.35, 0.9,
        [[("throughput  +57%", 15, GREEN, True)],
         [("고부하 tr 0.47 vs default 0.30 p/s · p95도 낮음", 11, INK, False)]])
rect(s, 7.15, 6.05, 5.6, 1.0, GRNL, line=GREEN, lw=1.2, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
textbox(s, 7.30, 6.12, 5.35, 0.9,
        [[("hit rate  0.67 유지 (default는 0.02 붕괴)", 15, GREEN, True)],
         [("tr 0.673 vs default 0.024 · 재프리필 ≈ 1/10", 11, INK, False)]])
notes(s,
"이제 결과입니다. 먼저 합성 워크로드, 동종 2×4090입니다. 이건 우리가 논문을 읽고 이렇게 나와야 한다고 "
"기대한 그림입니다. 왼쪽 throughput, 부하가 올라가면 tr이 default보다 57% 높습니다. 초당 0.47 대 0.30이고 "
"p95 지연도 tr이 더 낮습니다. 오른쪽 hit rate도 tr이 0.67을 유지하는데 default는 0.02로 붕괴합니다. 재프리필 "
"양은 tr이 default의 약 10분의 1입니다. 즉 tr이 프로그램 단위로 스래싱을 억제해서 캐시도 지키고 처리량도 "
"이깁니다. 논문의 핵심 주장을 우리 환경에서 그대로 재현한 거고, 역시 SOTA가 이긴다는 기대에 부합합니다. "
"그런데 다음 장에서 실데이터를 넣으면 이 그림이 뒤집힙니다.")
footer(s, "기대대로: tr이 스래싱 억제 → hit·throughput·p95 모두 우위 (합성, 2×4090)")

# ===========================================================================
# 9 — 결과 ② TraceLab (유지)
# ===========================================================================
s = slide()
title(s, "결과 ②  TraceLab 실데이터 · 2×4090", AMBER, "그런데 정반대 — 같은 tr인데 throughput을 잃는다 (hit은 여전히 지킴)")
add_image(s, "hetero_homo_tracelab_throughput.png", 1.70, width=5.7, left=0.75)
add_image(s, "hetero_homo_tracelab_hit_rate.png", 1.70, width=5.7, left=6.95)
rect(s, 0.90, 6.05, 5.6, 1.0, REDL, line=RED, lw=1.2, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
textbox(s, 1.05, 6.12, 5.35, 0.9,
        [[("throughput  −34%  (뒤집힘)", 15, RED, True)],
         [("고부하(c=48) tr 0.067 vs default 0.102 p/s", 11, INK, False)]])
rect(s, 7.15, 6.05, 5.6, 1.0, GRNL, line=GREEN, lw=1.2, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
textbox(s, 7.30, 6.12, 5.35, 0.9,
        [[("hit rate  0.77 유지 (여전히 압승)", 15, GREEN, True)],
         [("tr 0.772 vs default 0.026 — 약 30배 차이", 11, INK, False)]])
notes(s,
"그런데 같은 하드웨어, 같은 tr에 워크로드만 실데이터 TraceLab으로 바꾸면 그림이 뒤집힙니다. 왼쪽 throughput, "
"이번엔 tr이 default보다 34% 낮습니다. 0.067 대 0.102로 tr이 집니다. 방금 합성에서는 57% 이겼는데 정반대죠. "
"흥미로운 건 오른쪽 hit rate입니다. 여기서는 여전히 tr이 0.77을 유지하고 default는 0.03으로 붕괴합니다. 즉 "
"캐시는 똑같이 잘 지키는데 처리량만 잃습니다. 정리하면, tr은 어느 워크로드든 캐시는 확실히 지키지만, "
"throughput은 합성에서 이기고 실데이터에서 집니다. 여기서 오늘의 핵심 질문이 나옵니다 — 같은 스케줄러가 왜 "
"어떤 워크로드에선 이기고 어떤 워크로드에선 질까? 다음 장에서 이걸 하나의 식으로 설명하겠습니다.")
footer(s, "합성에선 +57%(승), TraceLab에선 −34%(패). hit은 둘 다 지킴 — 왜 승패가 갈릴까?")

# ===========================================================================
# 10 — 원인 ① R 모델 (도입부에 'compute 안 봄' 리드 추가)
# ===========================================================================
s = slide()
title(s, "왜 갈리나 — 원인 ①", RED, "tr은 메모리만 본다 → GPU 연산이 남아돈다 → 그걸 정량화한 게 R")
# ★ 리드 문장 (평서문, 먼저)
rect(s, 0.90, 1.48, 11.85, 0.56, REDL, line=RED, lw=1.3, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
textbox(s, 0.90, 1.48, 11.85, 0.56,
        [[("tr은 '메모리(용량=스래싱 회피)'만 보고 compute는 고려하지 않는다  →  GPU 연산이 남아돈다(놀고 있다).",
           13.5, RED, True)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
# 방정식 배너
rect(s, 0.90, 2.14, 11.85, 0.56, INK, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
textbox(s, 0.90, 2.14, 11.85, 0.56,
        [[("'얼마나 노는가'를 정량화:   U ≈ min(R, 1),   R = k_fit × d", 17, WHITE, True),
          ("   ( k_fit = 동시 적재 프로그램 수,  d = duty )", 11, LGRAY, False)]],
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
GBX, GBW = 3.55, 7.55
def gpu_gantt(y_label, y_res, y_bar, panel_fill, k_txt, res_runs, bar_kind, u_txt, cap_txt):
    rect(s, 0.80, y_label, 2.55, 0.98, panel_fill, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    textbox(s, 0.90, y_label+0.05, 2.4, 0.9, k_txt, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE)
    textbox(s, GBX, y_res, 9.2, 0.32, [res_runs])
    if bar_kind == "solid":
        rect(s, GBX, y_bar, GBW, 0.48, GREEN)
    else:
        rect(s, GBX, y_bar, GBW, 0.48, LGRAY)
        for cx in (0.0, 1.55, 3.10, 4.65, 6.20):
            rect(s, GBX+cx, y_bar, 0.62, 0.48, GREEN if bar_kind=="gappy_g" else RED)
    textbox(s, GBX+GBW+0.12, y_bar, 1.5, 0.48,
            [[(u_txt, 13.5, GREEN if bar_kind=="solid" else RED, True)]], anchor=MSO_ANCHOR.MIDDLE)
    textbox(s, GBX, y_bar+0.50, 9.2, 0.3, [[(cap_txt, 10, GRAY, False)]])
gpu_gantt(2.86, 2.74, 3.24, GRNL,
          [[("합성", 12.5, GREEN, True)], [("balanced", 10, INK, False)],
           [("d≈0.56 · k_fit≈3", 10, GRAY, False)]],
          [("R = 3 × 0.56 ≈ 1.7  ", 12.5, INK, True), ("≥ 1", 12.5, GREEN, True),
           ("   →   GPU 가득   →   ", 12, INK, False), ("tr 승 (+57%)", 12.5, GREEN, True)],
          "solid", "U ≈ 1.00",
          "GPU 타임라인: 동시 적재 3개가 서로의 툴콜 구간을 메워 GPU가 쉬지 않음")
gpu_gantt(4.62, 4.50, 5.00, REDL,
          [[("TraceLab", 12.5, RED, True)], [("prefill-heavy", 10, INK, False)],
           [("d≈0.18 · k_fit≈2", 10, GRAY, False)]],
          [("R = 2 × 0.18 ≈ 0.37  ", 12.5, INK, True), ("< 1", 12.5, RED, True),
           ("   →   GPU에 bubble   →   ", 12, INK, False), ("tr 패 (−34%)", 12.5, RED, True),
           ("  · 실측 U 36% 일치", 10.5, GRAY, False)],
          "gappy_r", "U ≈ 0.38",
          "GPU 타임라인: tr이 캐시 지키려 pause → 동시 적재 2개뿐 → 툴콜 구간이 그대로 GPU 구멍(bubble)")
rect(s, 0.90, 6.55, 11.85, 0.48, AMBRL, line=AMBER, lw=1.3, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
textbox(s, 0.90, 6.55, 11.85, 0.48,
        [[("차이는 워크로드 duty(d) 하나 — ", 12, INK, True),
          ("d가 R을, R이 GPU 활용률 U를, U가 승패를 결정 (같은 HW·같은 tr).", 12, AMBER, True)]],
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
notes(s,
"앞의 뒤집힘을 하나의 식으로 설명합니다. 이게 오늘 발표의 중심입니다. 먼저 가장 중요한 직관 한 줄 — 맨 위 "
"빨간 띠입니다. tr은 메모리, 즉 용량이 넘쳐서 스래싱이 나는지만 봅니다. 정작 GPU 연산 자원, compute는 보지 "
"않습니다. 그래서 캐시는 잘 지키지만 GPU 연산이 남아돌아 놀 수 있습니다. 그럼 그게 얼마나 노는지를 어떻게 "
"정량화하느냐 — 그게 R입니다. tr의 GPU 활용률 U는 R로 정해지고, R은 동시에 GPU에 올라간 프로그램 수 "
"k_fit 곱하기 워크로드 duty d입니다. duty는 프로그램이 전체 시간 중 실제로 GPU를 쓰는 비율입니다. 위쪽 "
"합성은 k_fit 3, d 0.56이라 R이 약 1.7로 1보다 큽니다. 세 프로그램이 서로의 툴콜 대기 구간을 번갈아 메우니 "
"GPU가 꽉 차고, U가 1에 가까워 tr이 이깁니다. 아래 TraceLab은 반대입니다. 입력이 길어 프로그램 하나가 KV를 "
"많이 잡으니 tr이 캐시를 지키려 pause를 해서 동시 적재가 2개로 줄고, d도 0.18로 작습니다. R이 0.37이라 "
"툴콜로 노는 구간을 다 못 메워 GPU에 구멍, bubble이 생기고 U가 0.38로 떨어져 집니다. 이 0.38이 실측 활용률 "
"36%와 거의 일치했습니다. 핵심은, 하드웨어도 스케줄러도 그대로인데 승패를 가른 건 워크로드 duty 하나이고, "
"그게 R을 통해 활용률을, 활용률이 승패를 결정한다는 겁니다. 그럼 tr은 왜 굳이 pause를 할까? 다음 장입니다.")
footer(s, "리드: tr은 메모리만 봄 → 연산 남아돎 → U≈min(R,1), R=k_fit×d · 합성 1.7(승)/TraceLab 0.37(패)")

# ===========================================================================
# 11 — 원인 ② 스케줄러 메커니즘 (본문 축소, 세부는 노트)
# ===========================================================================
s = slide()
title(s, "왜 갈리나 — 원인 ②", RED, "스케줄러는 '용량(스래싱 회피)'만 본다 — 두 병목이 코드에 있다")
textbox(s, 0.80, 1.48, 11.95, 0.34,
        [[("cost model = '용량 부등식'을 pause/resume/신규배정으로 유지. 판단 기준은 오직 KV 용량 — ", 11.5, GRAY, False),
          ("GPU 활용률(U)도 GPU별 용량도 목표가 아니다.", 11.5, RED, True)]])
# 좌: 용량 부등식 (간소화)
rect(s, 0.80, 2.00, 4.15, 3.55, AMBRL, line=AMBER, lw=1.6, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
textbox(s, 1.00, 2.12, 3.8, 3.35,
        [[("스케줄러가 보는 것 (매 5초)", 12.5, AMBER, True)],
         [("프로그램 = (컨텍스트 c, 상태)", 11, INK, False)],
         [("", 10, INK, False)],
         [("용량 부등식 = 스래싱 판정", 12, AMBER, True)],
         [("Σc − shared + buffer > C_total ?", 11, INK, True)],
         [("넘으면 pause · 여유면 resume", 10.5, INK, False)],
         [("", 12, INK, False)],
         [("→ 판단 기준은 오직 KV 용량.", 11.5, RED, True)],
         [("U도 GPU별 용량도 안 본다.", 11.5, RED, True)]])
arrow(s, 4.95, 3.1, 5.35, 3.1, RED, 2.5)
arrow(s, 4.95, 4.6, 5.35, 4.6, RED, 2.5)
# 우: 두 병목 콜아웃
rect(s, 5.40, 2.00, 7.35, 1.66, PANEL, line=RED, lw=1.7, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
textbox(s, 5.60, 2.10, 7.0, 1.5,
        [[("(a) 캐시 지키려 GPU를 굶긴다  →  R<1", 13, RED, True)],
         [("pause된 요청은 _wait_for_resume로 큐에서 블록 (router.py:416)", 11, INK, False)],
         [("→ resident(동시 적재)↓ → GPU idle↑ → throughput↓", 11, INK, True)],
         [("= 앞 장의 'R이 1보다 작아지는' 바로 그 메커니즘", 10.5, GRAY, False)]])
rect(s, 5.40, 3.86, 7.35, 1.69, PANEL, line=RED, lw=1.7, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
textbox(s, 5.60, 3.96, 7.0, 1.55,
        [[("(b) 큰 GPU를 덜 채운다  →  이종 확장의 걸림돌", 13, RED, True)],
         [("신규배정 = active_program_tokens 절대 최소 (router.py:358)", 11, INK, False)],
         [("C_total로 나누지 않음 = 용량 비례가 아님", 11, GRAY, False)],
         [("→ 2.03× 큰 5090에 2배 안 보냄 · split 0.5 ≪ 용량비 0.67", 11, RED, True)]])
# 하단 펀치라인
rect(s, 0.80, 5.72, 11.95, 0.62, INK, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
textbox(s, 0.80, 5.72, 11.95, 0.62,
        [[("판단 기준이 '용량'뿐이라 → (a) 캐시 지키려 GPU 굶김(R<1) · (b) 큰 GPU 덜 채움. 둘 다 코드 한두 줄이 근거.",
           12.5, WHITE, True)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
notes(s,
"앞 장에서 tr이 메모리만 보고 연산을 안 본다고 했는데, 그게 코드에서 실제로 어떻게 도는지입니다. 밀도를 "
"줄여 두 병목만 남겼고, 세부 알고리즘은 말로 보충하겠습니다. 왼쪽 — 스케줄러는 워크플로우 하나를 프로그램 "
"P라는 튜플로 봅니다. 컨텍스트 토큰 c, 지금 추론 중인지 툴 실행 중인지, 활성인지 대기인지 같은 상태를 "
"담습니다. 이건 요청 단위 엔진엔 없는 정보입니다. 매 5초마다 각 백엔드의 총 용량과 사용률, 각 프로그램의 c와 "
"상태를 모아서, Reasoning 프로그램의 컨텍스트 합에 Acting 프로그램은 시간 감쇠를 곱해 더하고, prefix 캐시로 "
"절약된 shared를 빼고 버퍼를 더한 값이 총 용량 C_total을 넘으면 스래싱으로 판정합니다. 넘으면 pause, 여유가 "
"생기면 resume, 요청이 오면 신규 배정을 합니다. 참고로 pause는 재계산 비용이 컨텍스트 길이의 제곱에 "
"비례하기 때문에 짧은 프로그램부터 빼고, resume은 대기 중인 프로그램부터 여유 노드에 넣습니다. 이 세부는 "
"넘어가고, 핵심은 판단 기준이 오직 KV 용량이라는 점입니다. 목표인 GPU 활용률도, 이종에서 GPU별 용량 "
"차이도 보지 않습니다. 그래서 오른쪽 두 병목이 생깁니다. (a) 캐시를 지키려 pause하면 요청이 큐에서 블록되고 "
"동시 적재가 줄어 GPU가 굶습니다. 이게 앞 장의 R이 1보다 작아지는 메커니즘입니다. (b) 신규 배정이 "
"active_program_tokens가 절대적으로 가장 작은 백엔드를 고르는데, 이 값을 용량으로 나누지 않아서 2배 큰 "
"5090에 2배를 안 보내고 과소활용합니다. 두 병목 모두 router.py의 특정 라인에 근거가 있고, 이게 향후 "
"과제로 이어집니다.")
footer(s, "판단=용량 부등식뿐 → (a) 캐시 지키려 GPU 굶김(R<1, :416) · (b) 큰 GPU 덜 채움(절대토큰, :358)")

# ===========================================================================
# 12 — 증명 ① 실험 C (인과, R 모델) (유지)
# ===========================================================================
s = slide()
title(s, "증명 ①  실험 C — 인과 확정 (완료)", GREEN, "앞의 '한 식'을 통제 실험으로 증명 — 예측 U vs 실측 U  r = 0.96")
add_image(s, "expC_R1_crossing.png", 1.50, width=11.5, left=0.90)
textbox(s, 0.90, 4.84, 11.5, 0.30,
        [[("duty 축으로 R을 0.31→1.0으로 밀자 U·throughput이 예측대로 회복 (R=1에서 교차)", 11, GRAY, False)]],
        align=PP_ALIGN.CENTER)
rect(s, 0.90, 5.18, 7.15, 1.72, PANEL, line=GREEN, lw=1.3, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
textbox(s, 1.10, 5.26, 6.85, 1.6,
        [[("R 모델:  U ≈ min(R, 1),   R = k_fit × d", 13, GREEN, True),
          ("     (예측 U vs 실측 U  r = 0.959)", 10.5, GRAY, False)],
         [("격리 실험 C 실측 — tr: R≈0.31<1 → U≈0.35 (GPU 65% 유휴) → throughput 열세", 11, RED, True)],
         [("인과 결정타: R을 1까지 밀면 throughput 0.068→0.144(×2.1),", 11.5, INK, True)],
         [("   hit 유지 → tr이 default를 두 축 모두 역전", 11.5, GREEN, True)]])
rect(s, 8.25, 5.18, 4.5, 1.72, AMBRL, line=AMBER, lw=1.2, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
textbox(s, 8.45, 5.26, 4.2, 1.6,
        [[("주의 — 두 종류의 회복", 12, AMBER, True)],
         [("pause만 완화:", 11, RED, True), (" thru↑ 지만", 10.5, INK, False)],
         [("   hit 0.79→0.65 (나쁜 회복)", 10.5, RED, False)],
         [("duty↑ / 용량↑:", 11, GREEN, True), (" thru↑ +", 10.5, INK, False)],
         [("   hit 유지 (좋은 회복)", 10.5, GREEN, True)]])
notes(s,
"이제 이게 진짜 원인이 맞다를 실험으로 증명한 부분입니다. 실험 C, 완료된 결과입니다. 핵심은 R이라는 하나의 "
"값입니다. tr의 GPU 활용률 U는 R과 거의 같은데, R은 동시에 올라간 프로그램 수 k_fit 곱하기 워크로드 duty "
"d입니다. 예측한 U와 실측 U의 상관이 0.959로 거의 완벽합니다. tr은 캐시를 지키려고 pause를 많이 해서 "
"k_fit이 낮게 유지되고, 그래서 R이 약 0.31로 1보다 작습니다. R이 1보다 작으면 GPU가 그만큼 놉니다. 여기서 "
"결정적인 실험 — 워크로드 duty를 조절해서 R을 1까지 밀어봤더니, throughput이 0.068에서 0.144로 두 배 넘게 "
"오르면서 default를 처리량과 hit rate 두 축 모두에서 역전했습니다. 즉 R이 1보다 작은 게 tr이 지는 원인이라는 "
"걸 인과적으로 확정한 겁니다. 오른쪽 아래 주의점 하나 — pause만 완화해도 처리량은 오르지만 그때는 hit "
"rate가 떨어집니다. 이건 나쁜 회복이고, 진짜 해법은 duty를 올리거나 용량을 키우는 좋은 회복입니다.")
footer(s, "R<1이 tr 열세의 원인 — duty로 R을 1까지 밀면 hit 지키며 throughput 역전 (r=0.96)")

# ===========================================================================
# 13 — 증명 ② SWE-bench (2차 실데이터) (유지)
# ===========================================================================
s = slide()
title(s, "증명 ②  SWE-bench — 2차 실데이터 (완료)", GREEN, "duty가 다른 워크로드에서 R 모델이 승패를 미리 맞힌다")
add_image(s, "swebench_throughput.png", 1.62, width=5.55, left=0.72)
add_image(s, "swebench_hitrate.png", 1.62, width=5.55, left=6.55)
rect(s, 0.90, 5.83, 11.85, 1.14, PANEL, line=GREEN, lw=1.3, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
textbox(s, 1.10, 5.90, 11.5, 1.04,
        [[("같은 2×4090 — duty(d) 하나가 R을, R이 승패를 결정 (세 워크로드 · 한 식)", 13, INK, True)],
         [("합성 d≈0.56 → R≈1.7 → 승 (+57%)    ·    ", 11.5, GREEN, True),
          ("TraceLab d≈0.18 → R<1 → 패 (−34%)    ·    ", 11.5, RED, True),
          ("SWE d≈1.0 → R≫1 → 승 (+78~84%)", 11.5, GREEN, True)],
         [("SWE는 스윕 전에 R로 예측 → 적중. 흩어진 승패가 하나의 값으로 통일.", 11, GRAY, True)]])
notes(s,
"두 번째 증명은 완전히 다른 실데이터인 SWE-bench입니다. 이것도 완료된 결과입니다. SWE-bench는 코딩 "
"에이전트라 생성이 길어서 duty d가 거의 1입니다. 아까 TraceLab은 d가 0.18이었죠. R 모델대로라면 d가 크니까 "
"R이 1보다 훨씬 커지고, 그러면 tr이 이겨야 합니다. 그리고 저는 이 예측을 스윕을 돌리기 전에 미리 적어놨고, "
"실제로 맞았습니다. 왼쪽 throughput, 부하가 높을 때 tr이 default보다 78%에서 84% 더 높습니다. 앞의 "
"TraceLab에서 34% 졌던 것과 정반대입니다. 오른쪽 hit rate는 여기서도 tr이 0.8 근처를 지키고 default는 "
"붕괴합니다. 아래 띠가 오늘 발표의 하이라이트입니다. 똑같은 2×4090 하드웨어인데, 한 워크로드에선 tr이 34% "
"지고 다른 워크로드에선 84% 이깁니다. 이 정반대 결과를 가르는 유일한 변수가 워크로드 duty 하나이고, 그걸 "
"R 모델이 미리 예측했습니다. 흩어져 있던 승패가 하나의 값으로 통일된 겁니다.")
footer(s, "같은 HW, duty만 다름: R 모델이 승패를 사전 예측 → 두 실데이터로 확증 (SWE +78~84% 적중)")

# ===========================================================================
# 14 — 향후 계획 (유지)
# ===========================================================================
s = slide()
title(s, "향후 계획", BLUE, "원인을 알았으니, 이종·분산 환경의 해법으로")
plans = [
    ("① Related work 추가 조사", AMBER,
     ["Continuum — 툴 콜 중 KV cache TTL 유지로 재프리필 회피",
      "DistServe — prefill/decode 분리, 그 외 이종 서빙 기법 폭넓게"]),
    ("② 용량 비례 라우팅", GREEN,
     ["신규 배정을 절대 토큰이 아닌 용량 비례로 (router.py:358 핵심)",
      "5090 과소활용 해소 → 이종에서 default 역전 목표"]),
    ("③ 이종 → 분산으로 확장", BLUE,
     ["R 모델을 이종/분산 클러스터의 라우팅 설계 원리로 일반화",
      "duty·용량을 함께 보는 스케줄링 정책 탐색"]),
]
y = 1.75
for h, col, items in plans:
    rect(s, 0.90, y, 11.85, 1.55, PANEL, line=col, lw=1.4, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    rect(s, 0.90, y, 0.14, 1.55, col)
    textbox(s, 1.20, y+0.12, 11.3, 0.45, [[(h, 14.5, col, True)]])
    textbox(s, 1.20, y+0.60, 11.3, 0.9,
            [[("• "+items[0], 12, INK, False)], [("• "+items[1], 12, INK, False)]])
    y += 1.72
notes(s,
"마지막으로 향후 계획입니다. 원인을 알았으니 이제 해법으로 갑니다. 첫째, related work를 더 조사하겠습니다. "
"특히 Continuum이라는 연구가 툴 콜 중에 KV 캐시를 TTL로 유지해서 재프리필을 피하는데, 제 문제의식과 정확히 "
"맞닿아 있어 깊게 볼 계획입니다. DistServe 같은 프리필/디코드 분리 기법도 함께 보겠습니다. 둘째, 가장 직접적인 "
"해법으로 용량 비례 라우팅을 시도하겠습니다. 앞서 5090 과소활용의 원인이 router.py 358번 라인의 절대 토큰 "
"배정이었으니, 이걸 용량 비례로 바꾸면 이종에서 default를 역전할 수 있을 것으로 봅니다. 셋째, 이걸 이종을 "
"넘어 분산 클러스터까지 확장하겠습니다. R 모델을 라우팅 설계의 일반 원리로 삼아, duty와 용량을 함께 고려하는 "
"정책을 만드는 게 목표입니다. 이상입니다. 감사합니다. 질문 받겠습니다.")
footer(s, "Continuum/DistServe 조사 · 용량 비례 라우팅 · 이종→분산 확장")

prs.save(OUT)
print("SAVED", OUT)
print("slides:", len(prs.slides._sldIdLst))
