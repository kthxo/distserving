#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""인턴 세미나 발표 덱 (10분) — 탑다운: 자기소개 → 주제 → 진행상황 → 향후계획.
스토리: 문제 → ThunderAgent baseline → hetero/워크로드 실험 → 결과(hit 승, thru 승패)
        → 원인(STP cost model ↔ 구현 간극) → 인과증명(실험 C·R모델, SWE 2차) → 향후.
모든 수치·코드라인은 logs/*.md 실측값만 사용(지어내기 금지).
Output: slides/2026-07-07_intern_seminar_yunuikang.pptx
Font: Calibri(라틴) + Malgun Gothic(한글 렌더링) — 시각적 통일은 Calibri.
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn
from PIL import Image

FIGS = "/home/yunuikang/yunuikang_work/distserving/figures"
OUT  = "/home/yunuikang/yunuikang_work/distserving/slides/2026-07-07_intern_seminar_yunuikang.pptx"
os.makedirs(os.path.dirname(OUT), exist_ok=True)

# ---- palette (절제된 톤) ----
INK   = RGBColor(0x1f, 0x24, 0x30)
BLUE  = RGBColor(0x27, 0x46, 0x90)
AMBER = RGBColor(0xE0, 0x8A, 0x1E)
GREEN = RGBColor(0x2E, 0x7D, 0x46)
RED   = RGBColor(0xC0, 0x39, 0x2B)
GRAY  = RGBColor(0x6B, 0x72, 0x80)
LGRAY = RGBColor(0xEC, 0xEE, 0xF2)
PANEL = RGBColor(0xF4, 0xF6, 0xFB)
BLUEL = RGBColor(0xE4, 0xEA, 0xF6)
AMBRL = RGBColor(0xFB, 0xF0, 0xDD)
GRNL  = RGBColor(0xE3, 0xF1, 0xE8)
REDL  = RGBColor(0xF8, 0xE4, 0xE1)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LATIN = "Calibri"
EA    = "Malgun Gothic"   # 한글 글리프 fallback (Calibri엔 한글 없음)

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]

def slide():
    return prs.slides.add_slide(BLANK)

def _set_font(run, size, color, bold=False, italic=False):
    run.font.size = Pt(size); run.font.color.rgb = color
    run.font.bold = bold; run.font.italic = italic
    run.font.name = LATIN
    # 한글 East-Asian typeface 명시 (fallback 깔끔하게)
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {}); rPr.append(el)
        el.set("typeface", EA if tag == "a:ea" else LATIN)

def textbox(s, l, t, w, h, lines, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, sa=3):
    tb = s.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = Pt(4); tf.margin_right = Pt(4); tf.margin_top = Pt(2); tf.margin_bottom = Pt(2)
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        for (txt, sz, col, bold) in (ln if isinstance(ln, list) else [ln]):
            r = p.add_run(); r.text = txt; _set_font(r, sz, col, bold)
        p.space_after = Pt(sa); p.space_before = Pt(0)
    return tb

def rect(s, l, t, w, h, fill, line=None, lw=1.0, shape=MSO_SHAPE.RECTANGLE, radius=None):
    sp = s.shapes.add_shape(shape, Inches(l), Inches(t), Inches(w), Inches(h))
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line; sp.line.width = Pt(lw)
    sp.shadow.inherit = False
    return sp

def boxtext(s, l, t, w, h, lines, fill, line=None, lw=1.0, align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    rect(s, l, t, w, h, fill, line, lw, shape=shape)
    textbox(s, l, t, w, h, lines, align=align, anchor=anchor)

def arrow(s, x1, y1, x2, y2, color=GRAY, lw=2.25):
    cn = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    cn.line.color.rgb = color; cn.line.width = Pt(lw)
    le = cn.line._get_or_add_ln()
    tail = le.makeelement(qn('a:tailEnd'), {'type': 'triangle', 'w': 'med', 'len': 'med'})
    le.append(tail)
    cn.shadow.inherit = False
    return cn

def title(s, tag, tag_color, head):
    rect(s, 0.55, 0.42, 0.14, 0.62, tag_color)
    textbox(s, 0.80, 0.34, 12.0, 0.66, [[(head, 25, INK, True)]], anchor=MSO_ANCHOR.MIDDLE)
    textbox(s, 0.80, 1.00, 12.0, 0.34, [[(tag, 12.5, tag_color, True)]])
    rect(s, 0.80, 1.40, 11.95, 0.02, tag_color)

def add_image(s, name, top, width=None, center=True, left=None, max_h=None):
    path = os.path.join(FIGS, name)
    iw, ih = Image.open(path).size
    if width is None: width = 11.7
    w = Inches(width); h = Emu(int(w * ih / iw))
    if max_h is not None and h > Inches(max_h):
        h = Inches(max_h); w = Emu(int(h * iw / ih))
    if left is None:
        left = (SW - w) / 2 if center else Inches(0.8)
    else:
        left = Inches(left)
    s.shapes.add_picture(path, left, Inches(top), width=w, height=h)
    return Inches(top) + h

def notes(s, text):
    s.notes_slide.notes_text_frame.text = text

def table(s, data, l, t, w, col_w, header_fill=BLUE, fs=12, hfs=12, row_h=0.42,
          zebra=True, cell_colors=None):
    rows, cols = len(data), len(data[0])
    gt = s.shapes.add_table(rows, cols, Inches(l), Inches(t), Inches(w), Inches(row_h*rows)).table
    gt.first_row = False; gt.horz_banding = False
    for ci, cw in enumerate(col_w):
        gt.columns[ci].width = Inches(cw)
    for ri in range(rows):
        gt.rows[ri].height = Inches(row_h)
        for ci in range(cols):
            cell = gt.cell(ri, ci)
            cell.margin_left = Pt(6); cell.margin_right = Pt(6)
            cell.margin_top = Pt(2); cell.margin_bottom = Pt(2)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            if ri == 0:
                cell.fill.solid(); cell.fill.fore_color.rgb = header_fill
            elif cell_colors and cell_colors.get((ri, ci)):
                cell.fill.solid(); cell.fill.fore_color.rgb = cell_colors[(ri, ci)]
            elif zebra and ri % 2 == 0:
                cell.fill.solid(); cell.fill.fore_color.rgb = PANEL
            else:
                cell.fill.solid(); cell.fill.fore_color.rgb = WHITE
            txt, col, bold = data[ri][ci]
            tf = cell.text_frame; tf.word_wrap = True
            p = tf.paragraphs[0]; p.alignment = PP_ALIGN.LEFT
            r = p.add_run(); r.text = txt
            _set_font(r, hfs if ri == 0 else fs, WHITE if ri == 0 else col, bold or ri == 0)
    return gt

def C(txt, col=INK, bold=False):
    return (txt, col, bold)

def footer(s, txt):
    textbox(s, 0.80, 7.06, 11.95, 0.34, [[(txt, 11, GRAY, False)]])

# ===========================================================================
# 1 — Title + 자기소개
# ===========================================================================
s = slide()
rect(s, 0, 0, 13.333, 7.5, WHITE)
rect(s, 0, 2.62, 13.333, 0.055, AMBER)
textbox(s, 0.9, 1.05, 11.6, 1.6,
        [[("이종(Heterogeneous) GPU 클러스터에서의", 30, BLUE, True)],
         [("에이전트 서빙 라우팅", 30, BLUE, True)]])
textbox(s, 0.9, 2.85, 11.6, 1.4,
        [[("SOTA 스케줄러 ", 15, GRAY, False), ("ThunderAgent", 15, AMBER, True),
          ("를 baseline으로, 이종 GPU·다양한 워크로드에서", 15, GRAY, False)],
         [("언제·왜 성능이 갈리는지를 하나의 인과 모델로 설명한다.", 15, GRAY, False)]])
textbox(s, 0.9, 6.25, 11.6, 0.9,
        [[("강윤의  (Yunui Kang)", 16, INK, True)],
         [("인턴 세미나  ·  개별연구 중간 발표  ·  2026-07-07", 12.5, GRAY, False)]])
notes(s,
"안녕하세요. 개별연구를 진행하고 있는 강윤의입니다. 오늘은 10분 동안 제 개별연구를 소개하겠습니다. "
"한 줄로 말씀드리면, 제 주제는 '서로 다른 종류의 GPU가 섞여 있는 클러스터에서, 들어오는 에이전트 요청을 "
"어느 GPU로 보낼지'를 연구하는 것입니다. 발표는 네 부분입니다. 먼저 제가 왜 이 문제를 푸는지, 다음으로 "
"지금까지 무엇을 했고 어떤 결과가 나왔는지, 마지막으로 앞으로 무엇을 할지 순서로 말씀드리겠습니다. "
"핵심 메시지 하나만 미리 말씀드리면 — 여러 실험 결과가 흩어져 있었는데, 그것들을 'R'이라는 하나의 값으로 "
"언제 이기고 언제 지는지 예측할 수 있게 됐다는 것입니다.")
footer(s, "개별연구 · 이종 GPU 에이전트 서빙")

# ===========================================================================
# 2 — 문제 정의: 왜 에이전트 서빙이 어려운가
# ===========================================================================
s = slide()
title(s, "문제 정의", RED, "에이전트 서빙은 왜 기존 LLM 서빙과 다른가")
textbox(s, 0.80, 1.55, 11.95, 0.5,
        [[("한 문장: ", 15, INK, True),
          ("에이전트는 '리즈닝 ↔ 툴 콜'을 여러 턴 반복 → KV 캐시가 계속 쌓이고, 이 캐시를 어느 GPU에 둘지가 성능을 지배한다.",
           15, INK, False)]])

# ReAct 루프 도식 (왼쪽)
lx, ly = 0.90, 2.35
boxtext(s, lx, ly, 2.05, 0.7, [[("REASONING", 12.5, WHITE, True)], [("GPU 추론", 10, BLUEL, False)]], BLUE)
boxtext(s, lx+3.05, ly, 2.05, 0.7, [[("ACTING", 12.5, WHITE, True)], [("툴 콜 (GPU idle)", 10, AMBRL, False)]], AMBER)
arrow(s, lx+2.05, ly+0.35, lx+3.05, ly+0.35, GRAY, 2.0)
arrow(s, lx+3.05, ly+0.62, lx+2.05, ly+0.62, GRAY, 2.0)
textbox(s, lx, ly+0.78, 5.1, 0.4, [[("1턴 = 리즈닝+툴콜, 멀티턴 반복 → KV 캐시 누적", 10.5, GRAY, False)]], align=PP_ALIGN.CENTER)

# 4가지 핵심 어려움 (오른쪽 카드)
cards = [
    ("① 멀티턴 KV 축적", "턴마다 캐시가 쌓임 → 요청 하나가 아니라 프로그램 전체를 관리해야", BLUE),
    ("② 툴 콜 버블", "툴 실행 중 GPU가 놀음(idle) → 이 빈 구간을 채우는 게 관건", AMBER),
    ("③ KV 캐시 스래싱", "GPU 자리 확보하려 남의 캐시를 evict → 다음 턴 전부 재프리필", RED),
    ("④ Locality ↔ Balancing", "캐시는 한 GPU에 묶여야(locality) vs 부하는 고루 나눠야(balancing) — 상충", GREEN),
]
cy = 2.35
for i, (h, d, col) in enumerate(cards):
    y = cy + i*1.05
    rect(s, 6.35, y, 6.35, 0.92, PANEL, line=col, lw=1.5, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    rect(s, 6.35, y, 0.12, 0.92, col)
    textbox(s, 6.62, y+0.06, 6.0, 0.42, [[(h, 13.5, col, True)]])
    textbox(s, 6.62, y+0.46, 6.0, 0.44, [[(d, 11, INK, False)]])
notes(s,
"먼저 문제 정의입니다. 기존 챗봇형 LLM은 입력 한 번에 출력 한 번이면 끝입니다. 그런데 에이전트는 다릅니다. "
"왼쪽 그림처럼 '리즈닝'과 '툴 콜'을 왔다갔다 반복합니다. 리즈닝은 GPU를 쓰지만, 툴 콜은 외부 API나 코드 "
"실행이라 그동안 GPU가 놉니다. 이게 여러 턴 반복되면서 네 가지 어려움이 생깁니다. "
"첫째, 턴마다 KV 캐시가 계속 쌓입니다. 그래서 요청 하나가 아니라 프로그램 전체를 통으로 관리해야 합니다. "
"둘째, 툴 콜 동안 GPU가 노는 '버블'이 생깁니다. 셋째, 여러 요청을 한 GPU에 올리다 보면 자리가 모자라 남의 "
"캐시를 지워버리고, 그 프로그램의 다음 턴이 오면 처음부터 다시 계산해야 하는 '스래싱'이 터집니다. 넷째, "
"캐시는 한 GPU에 묶여 있어야 재사용되는데(로컬리티), 부하는 여러 GPU에 고루 나눠야 합니다(밸런싱). 이 둘이 "
"정면으로 충돌합니다. 이 네 가지가 제 연구가 다루는 문제입니다.")
footer(s, "리즈닝↔툴콜 반복 → KV 축적·버블·스래싱·로컬리티/밸런싱 상충")

# ===========================================================================
# 3 — 연구 목표 & 접근
# ===========================================================================
s = slide()
title(s, "연구 목표 & 접근", BLUE, "기존 가정(동일 GPU)의 한계 → 이종 GPU로 확장")

boxtext(s, 0.90, 1.85, 5.55, 1.55,
        [[("기존 연구의 한계", 15, RED, True)],
         [("", 6, INK, False)],
         [("대부분 ", 12.5, INK, False), ("동일한(homogeneous) GPU", 12.5, RED, True),
          ("를 가정.", 12.5, INK, False)],
         [("→ GPU마다 다른 용량·처리량·대역폭을 고려하지 않음", 12, GRAY, False)]],
        REDL, line=RED, lw=1.5, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP)
arrow(s, 6.55, 2.62, 7.35, 2.62, GRAY, 2.5)
boxtext(s, 7.45, 1.85, 5.30, 1.55,
        [[("우리 연구의 목표", 15, GREEN, True)],
         [("", 6, INK, False)],
         [("서로 다른(heterogeneous) GPU", 12.5, GREEN, True),
          (" 클러스터로 확장", 12.5, INK, False)],
         [("→ 용량·처리량·대역폭 차이까지 고려한 라우팅", 12, GRAY, False)]],
        GRNL, line=GREEN, lw=1.5, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP)

boxtext(s, 0.90, 3.85, 11.85, 1.15,
        [[("접근 — 맨땅에서 시작하지 않는다: 이 문제의 SOTA를 baseline으로 삼아 병목을 찾는다.", 14.5, INK, True)],
         [("Related work 조사 → ", 12.5, GRAY, False), ("ThunderAgent (Kang et al., ICML 2026)", 12.5, AMBER, True),
          (" = 프로그램 단위 스케줄링으로 스래싱을 제어하는 SOTA", 12.5, GRAY, False)]],
        PANEL, line=AMBER, lw=1.5, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE)

textbox(s, 0.90, 5.25, 11.85, 1.4,
        [[("핵심 질문", 14, BLUE, True)],
         [("ThunderAgent를 ", 13, INK, False), ("이종 GPU + 다양한 워크로드", 13, BLUE, True),
          ("에 넣으면 어떤 병목이 드러나는가? — 이를 데이터로 규명하고, 원인을 알고리즘(코드)으로 설명한다.", 13, INK, False)]])
notes(s,
"그럼 목표와 접근입니다. 왼쪽 — 기존 연구는, 방금 소개한 SOTA를 포함해서, 대부분 '모든 GPU가 똑같다'고 "
"가정합니다. 하지만 현실 클러스터는 4090, 5090처럼 서로 다른 GPU가 섞여 있고, 그러면 용량도 처리량도 "
"메모리 대역폭도 다 다릅니다. 오른쪽 — 그래서 저는 이 차이까지 고려한 라우팅으로 문제를 확장하는 걸 "
"목표로 잡았습니다. 접근은, 맨땅에서 새 알고리즘을 만드는 게 아니라, 이 분야 SOTA인 ThunderAgent를 "
"baseline으로 두고, 그걸 이종 환경과 다양한 워크로드에 넣었을 때 '어디서 왜 막히는지'를 먼저 데이터로 "
"찾는 것입니다. 그리고 그 원인을 감이 아니라 실제 코드 로직으로 설명하는 데 초점을 뒀습니다.")
footer(s, "homogeneous 가정의 한계 → ThunderAgent(SOTA) baseline으로 이종 병목 탐색")

# ===========================================================================
# 4 — System Overview 다이어그램 (native shapes)
# ===========================================================================
s = slide()
title(s, "System Overview", BLUE, "Client → Router → 이종 인스턴스, 그 사이의 결정 요소")

# Client
boxtext(s, 0.75, 2.55, 1.7, 1.0, [[("Client", 14, WHITE, True)], [("에이전트\n프로그램", 10, BLUEL, False)]],
        INK, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
arrow(s, 2.45, 3.05, 3.35, 3.05, GRAY, 2.5)

# Router / Orchestrator
rect(s, 3.35, 2.05, 3.55, 2.6, PANEL, line=BLUE, lw=2.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
textbox(s, 3.35, 2.15, 3.55, 0.5, [[("Router / Orchestrator", 13.5, BLUE, True)]], align=PP_ALIGN.CENTER)
for i, (t, c) in enumerate([("① 신규 프로그램 라우팅", INK), ("② pause / resume", AMBER),
                            ("③ KV locality 유지", GREEN), ("④ 용량·부하 판단", INK)]):
    textbox(s, 3.5, 2.66+i*0.44, 3.3, 0.4, [[(t, 11.5, c, i in (0,1,2))]], align=PP_ALIGN.LEFT)

# arrows to instances
arrow(s, 6.9, 2.75, 8.15, 2.35, BLUE, 2.5)
arrow(s, 6.9, 3.55, 8.15, 4.35, BLUE, 2.5)
# pause/resume return arrow (dashed feel via amber)
arrow(s, 8.15, 3.05, 6.9, 3.05, AMBER, 1.75)
textbox(s, 6.95, 2.72, 1.25, 0.3, [[("route", 9.5, GRAY, False)]], align=PP_ALIGN.CENTER)
textbox(s, 6.95, 3.14, 1.25, 0.3, [[("pause↩", 9.5, AMBER, True)]], align=PP_ALIGN.CENTER)

# Instance 4090 (small)
rect(s, 8.20, 1.75, 4.55, 1.30, WHITE, line=GRAY, lw=1.5, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
textbox(s, 8.35, 1.83, 4.3, 0.4, [[("Instance A — RTX 4090", 13, INK, True)]])
rect(s, 8.35, 2.25, 1.55, 0.62, BLUEL, line=BLUE, lw=1.0)
textbox(s, 8.35, 2.28, 1.55, 0.56, [[("KV pool", 9.5, BLUE, True)], [("43,888 tok", 10.5, INK, True)]], align=PP_ALIGN.CENTER)
textbox(s, 10.05, 2.24, 2.6, 0.7, [[("작은 용량 · 낮은 대역폭", 10.5, GRAY, False)],
                                     [("resident ≈ 2 프로그램", 10.5, GRAY, False)]])

# Instance 5090 (large)
rect(s, 8.20, 3.55, 4.55, 1.55, WHITE, line=GREEN, lw=1.8, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
textbox(s, 8.35, 3.63, 4.3, 0.4, [[("Instance B — RTX 5090", 13, INK, True)]])
rect(s, 8.35, 4.05, 2.05, 0.72, GRNL, line=GREEN, lw=1.0)
textbox(s, 8.35, 4.10, 2.05, 0.64, [[("KV pool", 9.5, GREEN, True)], [("89,040 tok", 11, INK, True)]], align=PP_ALIGN.CENTER)
textbox(s, 10.55, 4.02, 2.15, 0.9, [[("큰 용량(2.03×)", 10.5, GREEN, True)],
                                      [("높은 처리량/대역폭", 10.5, GRAY, False)],
                                      [("resident ≈ 4", 10.5, GRAY, False)]])

# 상태기계 (하단 좌)
rect(s, 0.90, 5.35, 6.0, 1.55, PANEL, line=LGRAY, lw=1.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
textbox(s, 1.05, 5.42, 5.7, 0.35, [[("요청 생애주기 상태기계", 12, INK, True)]])
boxtext(s, 1.05, 5.85, 1.55, 0.62, [[("ACTIVE", 11.5, WHITE, True)]], GREEN)
boxtext(s, 3.05, 5.85, 1.55, 0.62, [[("PAUSED", 11.5, WHITE, True)]], AMBER)
boxtext(s, 5.05, 5.85, 1.55, 0.62, [[("resume", 11.5, WHITE, True)]], BLUE)
arrow(s, 2.60, 6.16, 3.05, 6.16, GRAY, 2.0)
arrow(s, 4.60, 6.16, 5.05, 6.16, GRAY, 2.0)
textbox(s, 1.05, 6.50, 5.7, 0.35, [[("Eq.6(용량) 위반 시 pause → 프록시 큐에서 대기 → 여유 생기면 resume", 10, GRAY, False)]])

# 팩터 요약 (하단 우)
rect(s, 7.10, 5.35, 5.65, 1.55, AMBRL, line=AMBER, lw=1.2, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
textbox(s, 7.25, 5.42, 5.4, 1.45,
        [[("Router가 저울질하는 상충 요소", 12, AMBER, True)],
         [("• KV locality: 프로그램을 같은 GPU에 고정(캐시 재사용)", 10.5, INK, False)],
         [("• Load balancing: 부하를 이종 용량비대로 분산", 10.5, INK, False)],
         [("• 이 둘의 균형이 곧 라우팅 정책 = 오늘 분석 대상", 10.5, INK, True)]])
notes(s,
"교수님이 그림을 좋아하셔서 시스템 구조를 한 장에 정리했습니다. 왼쪽부터, 클라이언트가 에이전트 프로그램을 "
"보내면 가운데 Router가 받습니다. Router가 하는 결정이 네 가지입니다. 새 프로그램을 어느 GPU에 보낼지(신규 "
"라우팅), 메모리가 부족하면 어떤 프로그램을 잠시 빼고(pause) 나중에 다시 넣을지(resume), 같은 프로그램의 "
"다음 턴을 캐시가 있는 같은 GPU로 보낼지(로컬리티), 그리고 각 GPU의 용량과 부하를 어떻게 볼지입니다. "
"오른쪽이 인스턴스인데, 여기가 제 연구의 핵심입니다 — 4090은 KV 풀이 약 4만4천 토큰이라 프로그램 두 개 "
"정도만 동시에 올라가고, 5090은 약 8만9천 토큰으로 2배 크고 처리량도 높습니다. 아래 왼쪽은 요청의 상태 "
"변화입니다. 메모리 용량 조건(논문 Eq.6)을 위반하면 프로그램이 PAUSED 상태로 프록시 큐에서 대기하다가, "
"자리가 나면 resume됩니다. 아래 오른쪽 — 결국 Router는 캐시를 위한 로컬리티와 부하 분산이라는 상충하는 "
"두 요소를 저울질하는데, 이 저울질 방식이 오늘 분석의 대상입니다.")
footer(s, "이종 인스턴스(4090 작음 / 5090 2.03× 큼) · Router가 라우팅·pause/resume·locality를 결정")

# ===========================================================================
# 5 — 실험 세팅
# ===========================================================================
s = slide()
title(s, "실험 세팅", BLUE, "동종 2×4090에서 워크로드(duty)만 바꿔가며 tr vs default")

boxtext(s, 0.90, 1.75, 3.75, 1.65,
        [[("무엇을 비교?", 13, BLUE, True)],
         [("tr", 13, GREEN, True), ("  = ThunderAgent 스케줄러", 11.5, INK, False)],
         [("     (프로그램 단위 pause/resume)", 10, GRAY, False)],
         [("default", 13, RED, True), (" = 최소부하로 프록시만", 11.5, INK, False)],
         [("     (프로그램 통제 없음)", 10, GRAY, False)]],
        PANEL, line=BLUE, lw=1.3, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP)

boxtext(s, 4.85, 1.75, 3.75, 1.65,
        [[("하드웨어", 13, BLUE, True)],
         [("2 × RTX 4090 (동종)", 11.5, INK, True)],
         [("KV 풀 43,888 tok / GPU", 10.5, GRAY, False)],
         [("이종(4090+5090) 용량차이는", 10.5, GRAY, False)],
         [("   메커니즘 분석·향후 과제로", 10.5, GRAY, False)]],
        PANEL, line=BLUE, lw=1.3, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP)

boxtext(s, 8.80, 1.75, 3.95, 1.65,
        [[("워크로드 — duty(d)로 대비", 13, BLUE, True)],
         [("합성 — balanced  d≈0.56", 11, INK, False)],
         [("TraceLab — prefill-heavy  d≈0.18", 11, INK, False)],
         [("SWE-bench — decode-heavy  d≈1.0", 11, INK, True)],
         [("d = GPU 실작동 시간 비율", 10, GRAY, False)]],
        PANEL, line=BLUE, lw=1.3, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP)

data = [
 [C("실험", WHITE, True), C("워크로드 (성격, duty d)", WHITE, True), C("HW", WHITE, True), C("이 실험이 답하는 것", WHITE, True)],
 [C("합성", GREEN, True), C("balanced (d≈0.56)"), C("2×4090"), C("tr이 의도대로 작동하는가 (기대)")],
 [C("D", BLUE, True), C("TraceLab · prefill-heavy (d≈0.18)"), C("2×4090"), C("실데이터에서도 그런가")],
 [C("C", GREEN, True), C("TraceLab (duty 스윕)"), C("격리 2×4090"), C("★ 승패 원인의 인과 증명 (R 모델)")],
 [C("SWE", GREEN, True), C("SWE-bench · decode-heavy (d≈1.0)"), C("2×4090"), C("★ 2차 실데이터로 결론 재현")],
]
table(s, data, 0.90, 3.80, 11.85, [1.0, 4.9, 1.3, 4.65], fs=12, hfs=12, row_h=0.54)
textbox(s, 0.90, 6.60, 11.85, 0.34,
        [[("모두 동종 2×4090 — 워크로드 duty만 바꾼다. 이종(5090) 용량차이 문제는 메커니즘(코드)·향후 과제에서 다룸.",
           10.5, GRAY, True)]])
notes(s,
"실험 세팅입니다. 비교하는 두 가지는, tr — ThunderAgent 스케줄러로 프로그램 단위로 pause와 resume을 하는 "
"방식이고, default — 부하 적은 GPU로 요청만 넘기고 프로그램 통제는 안 하는 방식입니다. 오늘 실험은 하드웨어를 "
"동종 2×4090으로 고정하고, 워크로드만 바꿔가며 봤습니다. 이렇게 하면 승패를 가르는 변수를 하드웨어가 아니라 "
"워크로드 하나로 깔끔하게 분리할 수 있습니다. 워크로드는 duty, 즉 전체 시간 중 GPU가 실제로 도는 비율로 "
"구분합니다. 합성은 중간(0.56), TraceLab은 입력이 길어 프리필이 무거운데 duty는 낮고(0.18), SWE-bench는 코딩 "
"에이전트라 생성이 길어 duty가 거의 1입니다. 표를 보시면, 먼저 합성으로 'tr이 기대대로 작동하는지' 확인하고, "
"TraceLab 실데이터로 '실제로도 그런지' 봅니다. 그다음 실험 C와 SWE가 오른쪽 별표, 오늘 강조할 인과 증명입니다. "
"이종 4090+5090의 용량 차이 문제는 뒤 메커니즘 분석에서 코드로 짚고, 향후 과제로 연결하겠습니다.")
footer(s, "동종 2×4090 고정 · 워크로드 duty(d)만 변화 · 합성→TraceLab→(원인)→SWE")

# ===========================================================================
# 6 — 결과 ① 합성 워크로드 (의도대로 작동)
# ===========================================================================
s = slide()
title(s, "결과 ①  합성 워크로드 · 2×4090", GREEN, "우리 기대대로 — tr이 스래싱을 억제해 throughput·hit 모두 이긴다")
add_image(s, "thrash_throughput.png", 1.70, width=5.7, left=0.75)
add_image(s, "thrash_hit_rate.png", 1.70, width=5.7, left=6.95)
# key figures band
rect(s, 0.90, 6.05, 5.6, 1.0, GRNL, line=GREEN, lw=1.2, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
textbox(s, 1.05, 6.12, 5.35, 0.9,
        [[("throughput  +57%", 15, GREEN, True)],
         [("고부하 tr 0.47 vs default 0.30 p/s · p95도 낮음", 11, INK, False)]])
rect(s, 7.15, 6.05, 5.6, 1.0, GRNL, line=GREEN, lw=1.2, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
textbox(s, 7.30, 6.12, 5.35, 0.9,
        [[("hit rate  0.67 유지 (default는 0.02 붕괴)", 15, GREEN, True)],
         [("tr 0.673 vs default 0.024 · 재프리필 ≈ 1/10", 11, INK, False)]])
notes(s,
"이제 결과입니다. 먼저 합성 워크로드, 동종 2×4090입니다. 이건 우리가 논문을 읽고 '이렇게 나와야 한다'고 "
"기대한 그림입니다. 두 축을 보시면 — 왼쪽 throughput, 부하가 올라가면 tr이 default보다 57% 높습니다. 초당 "
"0.47 대 0.30이고, p95 지연도 tr이 더 낮습니다. 오른쪽 hit rate도 tr이 0.67을 유지하는데 default는 0.02로 "
"붕괴합니다. 재프리필 양은 tr이 default의 약 10분의 1입니다. 즉 tr이 프로그램 단위로 스래싱을 억제해서 캐시도 "
"지키고 처리량도 이깁니다. 논문의 핵심 주장을 우리 환경에서 그대로 재현한 거고, '역시 SOTA가 이긴다'는 "
"기대에 부합합니다. 그런데 — 다음 장에서 실데이터를 넣으면 이 그림이 뒤집힙니다.")
footer(s, "기대대로: tr이 스래싱 억제 → hit·throughput·p95 모두 우위 (합성, 2×4090)")

# ===========================================================================
# 7 — 결과 ② TraceLab 실데이터 (반대로 → 왜?)
# ===========================================================================
s = slide()
title(s, "결과 ②  TraceLab 실데이터 · 2×4090", AMBER, "그런데 정반대 — 같은 tr인데 throughput을 잃는다 (hit은 여전히 지킴)")
add_image(s, "hetero_homo_tracelab_throughput.png", 1.70, width=5.7, left=0.75)
add_image(s, "hetero_homo_tracelab_hit_rate.png", 1.70, width=5.7, left=6.95)
# key figures band
rect(s, 0.90, 6.05, 5.6, 1.0, REDL, line=RED, lw=1.2, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
textbox(s, 1.05, 6.12, 5.35, 0.9,
        [[("throughput  −34%  (뒤집힘)", 15, RED, True)],
         [("고부하(c=48) tr 0.067 vs default 0.102 p/s", 11, INK, False)]])
rect(s, 7.15, 6.05, 5.6, 1.0, GRNL, line=GREEN, lw=1.2, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
textbox(s, 7.30, 6.12, 5.35, 0.9,
        [[("hit rate  0.77 유지 (여전히 압승)", 15, GREEN, True)],
         [("tr 0.772 vs default 0.026 — 약 30배 차이", 11, INK, False)]])
notes(s,
"그런데 같은 하드웨어, 같은 tr에 워크로드만 실데이터 TraceLab으로 바꾸면 그림이 뒤집힙니다. 왼쪽 throughput — "
"이번엔 tr이 default보다 34% 낮습니다. 0.067 대 0.102로 tr이 집니다. 방금 합성에서는 57% 이겼는데 정반대죠. "
"흥미로운 건 오른쪽 hit rate입니다 — 여기서는 여전히 tr이 0.77을 유지하고 default는 0.03으로 붕괴합니다. "
"즉 캐시는 똑같이 잘 지키는데, 처리량만 잃습니다. 정리하면, tr은 어느 워크로드든 캐시는 확실히 지키지만, "
"throughput은 합성에서 이기고 실데이터에서 집니다. 여기서 오늘의 핵심 질문이 나옵니다 — 같은 스케줄러가 왜 "
"어떤 워크로드에선 이기고 어떤 워크로드에선 질까? 다음 장에서 이걸 하나의 식으로 설명하겠습니다.")
footer(s, "합성에선 +57%(승), TraceLab에선 −34%(패). hit은 둘 다 지킴 — 왜 승패가 갈릴까?")

# ===========================================================================
# 8 — 왜 갈리나: R 모델 (한 식) — GPU 활용률이 승패를 결정
# ===========================================================================
s = slide()
title(s, "왜 갈리나 — 원인 ①", RED, "GPU 활용률(U)이 승패를 결정한다")
# 방정식 배너
rect(s, 0.90, 1.55, 11.85, 0.60, INK, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
textbox(s, 0.90, 1.55, 11.85, 0.60,
        [[("U ≈ min(R, 1)          R = k_fit × d", 18, WHITE, True),
          ("        ( k_fit = 동시 적재 프로그램 수,   d = 워크로드 duty )", 11.5, LGRAY, False)]],
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

GBX, GBW = 3.55, 7.55   # GPU 바 x/width

def gpu_gantt(y_label, y_res, y_bar, panel_fill, k_txt, res_runs, bar_kind, u_txt, cap_txt):
    # 좌측 라벨
    rect(s, 0.80, y_label, 2.55, 1.0, panel_fill, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    textbox(s, 0.90, y_label+0.06, 2.4, 0.9, k_txt, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE)
    # 결과 라인
    textbox(s, GBX, y_res, 9.2, 0.34, [res_runs])
    # GPU 바 (라벨은 캡션으로 대체)
    if bar_kind == "solid":
        rect(s, GBX, y_bar, GBW, 0.50, GREEN)
    else:  # gappy: 회색 트랙 + 초록 청크
        rect(s, GBX, y_bar, GBW, 0.50, LGRAY)
        for cx in (0.0, 1.55, 3.10, 4.65, 6.20):
            rect(s, GBX+cx, y_bar, 0.62, 0.50, GREEN if bar_kind=="gappy_g" else RED)
    textbox(s, GBX+GBW+0.12, y_bar, 1.5, 0.50,
            [[(u_txt, 14, GREEN if bar_kind=="solid" else RED, True)]], anchor=MSO_ANCHOR.MIDDLE)
    # 캡션
    textbox(s, GBX, y_bar+0.54, 9.2, 0.3, [[(cap_txt, 10.5, GRAY, False)]])

# 합성 (R≥1, 승)
gpu_gantt(2.55, 2.40, 2.98, GRNL,
          [[("합성", 13, GREEN, True)], [("balanced", 10.5, INK, False)],
           [("d ≈ 0.56 · k_fit ≈ 3", 10.5, GRAY, False)]],
          [("R = 3 × 0.56 ≈ 1.7  ", 13, INK, True), ("≥ 1", 13, GREEN, True),
           ("   →   GPU 가득   →   ", 12.5, INK, False), ("tr 승 (+57%)", 13, GREEN, True)],
          "solid", "U ≈ 1.00",
          "GPU 타임라인: 동시 적재 3개가 서로의 툴콜 구간을 메워 GPU가 쉬지 않음")

# TraceLab (R<1, 패)
gpu_gantt(4.55, 4.40, 4.98, REDL,
          [[("TraceLab", 13, RED, True)], [("prefill-heavy", 10.5, INK, False)],
           [("d ≈ 0.18 · k_fit ≈ 2", 10.5, GRAY, False)]],
          [("R = 2 × 0.18 ≈ 0.37  ", 13, INK, True), ("< 1", 13, RED, True),
           ("   →   GPU에 bubble   →   ", 12.5, INK, False), ("tr 패 (−34%)", 13, RED, True),
           ("   ·  실측 U 36% 일치", 11, GRAY, False)],
          "gappy_r", "U ≈ 0.38",
          "GPU 타임라인: tr이 캐시 지키려 pause → 동시 적재 2개뿐 → 툴콜 구간이 그대로 GPU 구멍(bubble)")

# 펀치라인
rect(s, 0.90, 6.52, 11.85, 0.52, AMBRL, line=AMBER, lw=1.3, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
textbox(s, 0.90, 6.52, 11.85, 0.52,
        [[("차이는 워크로드 duty(d) 하나 — ", 12.5, INK, True),
          ("d가 R을, R이 GPU 활용률 U를, U가 승패를 결정한다 (같은 HW·같은 tr).", 12.5, AMBER, True)]],
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
notes(s,
"앞의 뒤집힘을 하나의 식으로 설명합니다. 이게 오늘 발표의 중심입니다. tr의 GPU 활용률 U는 R이라는 값으로 "
"정해지고, R은 동시에 GPU에 올라간 프로그램 수 k_fit 곱하기 워크로드 duty d입니다. 그리고 처리량은 결국 GPU를 "
"얼마나 바쁘게 쓰느냐, 즉 U가 결정합니다. 위쪽 합성을 보시면 — k_fit이 3, d가 0.56이라 R이 약 1.7로 1보다 "
"큽니다. 프로그램 세 개가 서로의 툴콜 대기 구간을 번갈아 메우니까 GPU가 쉬지 않고 꽉 찹니다. U가 1에 "
"가깝고, 그래서 tr이 이깁니다. 아래 TraceLab은 반대입니다 — 입력이 길어 프로그램 하나가 KV를 많이 잡으니 "
"tr이 캐시를 지키려고 pause를 해서 동시 적재가 2개로 줄고, d도 0.18로 작습니다. R이 0.37로 1보다 작으니, "
"프로그램들이 툴콜로 노는 구간을 다 못 메워서 GPU에 구멍, 즉 bubble이 생깁니다. U가 0.38로 떨어지고 tr이 "
"집니다. 그리고 이 0.38은 실측 GPU 활용률 36%와 거의 일치합니다. 핵심은 — 하드웨어도 스케줄러도 그대로인데, "
"승패를 가른 건 워크로드 duty 하나이고, 그게 R을 통해 GPU 활용률을, 활용률이 승패를 결정한다는 겁니다. "
"그럼 자연스러운 다음 질문 — tr은 왜 굳이 pause를 해서 동시 적재를 낮출까? 다음 장 메커니즘에서 답합니다.")
footer(s, "한 식: U ≈ min(R,1), R = k_fit×d · 합성 R≈1.7(승) / TraceLab R≈0.37(패) — duty가 승패를 가른다")

# ===========================================================================
# 9 — 스케줄러 메커니즘 다이어그램 (그림1 재현·구체화, native shapes)
# ===========================================================================
s = slide()
title(s, "왜 갈리나 — 원인 ②", RED,
      "스케줄러는 '용량(스래싱 회피)'만 본다 — 목표인 GPU 활용률·용량차이는 반영 안 함")
textbox(s, 0.75, 1.48, 11.95, 0.34,
        [[("cost model = 아래 '용량 부등식'을 pause/resume/신규배정으로 유지. 판단 기준은 오직 KV 용량. ", 11, GRAY, False),
          ("GPU 활용률(U)도 GPU별 용량(C_total)도 목표가 아니다.", 11, RED, True)]])

CX1, CW1 = 0.75, 3.55
CX2, CW2 = 4.45, 3.85
CX3, CW3 = 8.75, 4.00

# --- COL1: ① 수집 → ② 용량 부등식 ---
rect(s, CX1, 1.92, CW1, 1.30, LGRAY, line=GRAY, lw=1.2, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
textbox(s, CX1+0.18, 1.99, CW1-0.32, 1.18,
        [[("① 수집  (매 5초 loop)", 12, INK, True)],
         [("백엔드별: C_total · KV 사용률(kv_usage)", 10, INK, False)],
         [("프로그램별: c(컨텍스트) · τ(추론/행동) · s(상태)", 10, INK, False)],
         [("prefix 캐시 조회/히트 → shared(절약 토큰)", 10, GRAY, False)]])
arrow(s, CX1+CW1/2, 3.24, CX1+CW1/2, 3.53, GRAY, 2.5)
rect(s, CX1, 3.55, CW1, 2.12, AMBRL, line=AMBER, lw=1.6, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
textbox(s, CX1+0.18, 3.62, CW1-0.32, 2.0,
        [[("② 용량 부등식 — 스래싱 판정", 12, AMBER, True)],
         [("Σ c (Reasoning)  +  Σ c·f(t) (Acting)", 10.5, INK, True)],
         [("       − shared  +  buffer   >   C_total ?", 10.5, INK, True)],
         [("f(t)=2⁻ᵗ : 오래 논 Acting 캐시는 할인(축출 우선)", 9.5, GRAY, False)],
         [("shared=prefix 절약 · buffer=프로그램당 100", 9.5, GRAY, False)],
         [("논문: λmax / λmin 워터마크로 히스테리시스", 9.5, GRAY, False)],
         [("state.py:185–194  ·  Eq.6 운영판", 9, GRAY, False)]])

# --- COL2: 세 정책 (구체적 기준) ---
rect(s, CX2, 1.92, CW2, 1.16, BLUEL, line=BLUE, lw=1.3, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
textbox(s, CX2+0.18, 1.99, CW2-0.32, 1.04,
        [[("resume (넣기)", 12, BLUE, True), ("  — 여유 시  (사용 < λmin·C)", 9.5, GRAY, False)],
         [("REASONING 우선   Sʀ = 1/c + [τ=R]", 10, INK, False)],
         [("여유 노드에 BFD 배치 · router.py:807", 9.5, GRAY, False)]])
rect(s, CX2, 3.24, CW2, 1.28, REDL, line=RED, lw=1.3, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
textbox(s, CX2+0.18, 3.31, CW2-0.32, 1.16,
        [[("pause (빼기)", 12, RED, True), ("  — 초과 시  (사용 > λmax·C)", 9.5, GRAY, False)],
         [("ACTING 우선 · 짧은 c 부터   Sᴘ = 1/c + [τ=A]", 10, INK, False)],
         [("이유: recompute ∝ c²  (작은 것 여럿이 더 쌈)", 9.5, GRAY, False)],
         [("router.py:773–805", 9.5, GRAY, False)]])
rect(s, CX2, 4.68, CW2, 1.24, AMBRL, line=AMBER, lw=1.4, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
textbox(s, CX2+0.18, 4.75, CW2-0.32, 1.12,
        [[("신규 배정", 12, INK, True), ("  — 요청 도착 시", 9.5, GRAY, False)],
         [("후보 = remaining_capacity ≥ 필요량 인 백엔드", 10, INK, False)],
         [("선택 = active_program_tokens 절대 최소", 10, RED, True)],
         [("router.py:358", 9.5, GRAY, False)]])

# 분기 화살표 (② 부등식 → 세 정책)
arrow(s, CX1+CW1, 3.85, CX2, 2.55, GREEN, 2.3)   # → resume
arrow(s, CX1+CW1, 4.30, CX2, 3.90, RED, 2.3)     # → pause
arrow(s, CX1+CW1, 4.95, CX2, 5.28, GRAY, 2.3)    # → 신규 배정

# --- COL3: ★ 핵심 콜아웃 (5090 과소활용) ---
arrow(s, CX2+CW2, 5.30, CX3, 4.55, AMBER, 2.5)   # 신규배정 → 핵심
rect(s, CX3, 3.10, CW3, 2.55, AMBRL, line=AMBER, lw=1.9, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
textbox(s, CX3+0.22, 3.24, CW3-0.42, 2.35,
        [[("(b) 큰 GPU를 덜 채운다", 12.5, RED, True)],
         [("신규배정 기준 = active_program_tokens 절대 최소", 10.5, INK, False)],
         [("( C_total 로 나누지 않음 = 용량 비례 아님 )", 10.5, GRAY, False)],
         [("", 4, INK, False)],
         [("→ 용량 2.03× 큰 5090에 2배 안 보냄", 11.5, INK, True)],
         [("→ 실측 split ≈ 0.5 ≪ 용량비 0.67 (과소활용)", 11, RED, True)],
         [("→ 이종 확장의 걸림돌 (router.py:358, 향후 과제)", 10.5, RED, True)]])

# --- 하단: pause 병목 사슬 (throughput) ---
rect(s, CX1, 6.02, 7.55, 0.90, PANEL, line=BLUE, lw=1.2, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
textbox(s, CX1+0.22, 6.10, 7.15, 0.78,
        [[("(a) 캐시 지키려 GPU를 굶긴다: ", 11, BLUE, True),
          ("pause된 요청은 _wait_for_resume로 큐 블록 (router.py:416)", 10.5, INK, False)],
         [("→ resident(동시 적재) 수↓ → GPU idle↑ → throughput↓  (= R<1)", 10.5, RED, True)]])
notes(s,
"앞 장의 간극이 실제로 어떻게 도는지, 이번엔 무엇을 수집하고 무엇을 판정하는지까지 구체적으로 담았습니다. "
"먼저 스케줄러는 워크플로우 하나를 프로그램 P라는 튜플로 봅니다 — 컨텍스트 토큰 c, 지금 추론 중인지 툴 실행 "
"중인지(τ), 활성인지 대기인지(s) 같은 상태를 담습니다. 이 정보는 요청 단위 엔진엔 없는 것입니다. "
"① 수집 — 매 5초마다 각 백엔드의 총 KV 용량 C_total과 사용률, 각 프로그램의 c·τ·상태, 그리고 prefix 캐시로 "
"절약된 shared 토큰을 모읍니다. ② 용량 부등식 — 이걸로 스래싱을 판정합니다. Reasoning 프로그램의 컨텍스트는 "
"그대로 더하고, Acting 프로그램의 컨텍스트에는 시간 감쇠 f(t)=2의 −t승을 곱합니다. 오래 툴을 기다린 캐시는 "
"곧 비워질 테니 덜 세게 보는 겁니다. 거기서 prefix 절약분 shared를 빼고 프로그램당 버퍼를 더한 값이 C_total을 "
"넘으면 스래싱입니다. 논문은 여기에 고·저 워터마크 λmax·λmin으로 히스테리시스를 둡니다. "
"이 판정으로 세 가지 행동을 합니다. resume는 여유가 생기면(사용량이 λmin 아래) REASONING 프로그램부터, "
"점수 1/c + τ가 R이면 가점, 여유 노드에 BFD로 넣습니다. pause는 초과 시 ACTING부터, 그리고 컨텍스트가 짧은 "
"것부터 뺍니다 — 재계산 비용이 컨텍스트 길이의 제곱에 비례하기 때문에 작은 걸 여러 개 빼는 게 항상 쌉니다. "
"신규 배정은 요청이 오면, 여유가 충분한 후보 중에서 active_program_tokens가 절대적으로 가장 작은 백엔드를 "
"고릅니다. 오른쪽 별표가 핵심입니다 — 이 절대값 기준이 용량 C_total로 나누질 않기 때문에, 2배 큰 5090에 2배를 "
"안 보내고, 실측 split이 0.5로 용량비 0.67에 못 미쳐 5090을 과소활용합니다. 이게 (b) 큰 GPU를 덜 채우는 "
"문제입니다. 아래 파란 띠가 (a)입니다 — 캐시를 지키려 pause하면 요청이 큐에서 블록되고 동시 적재가 줄어 "
"GPU가 굶습니다. 이게 바로 앞 장의 R이 1보다 작아지는 메커니즘입니다. 핵심은 — 스케줄러의 유일한 판단 "
"기준이 '용량 부등식(스래싱 회피)'뿐이고, 정작 목표인 GPU 활용률도, 이종에서의 GPU별 용량 차이도 반영하지 "
"않습니다. 그래서 캐시를 지키려다 GPU를 굶기거나(a), 큰 GPU를 덜 채웁니다(b). 두 병목 모두 코드 한두 줄에 "
"근거가 있습니다.")
footer(s, "판단 기준 = 용량 부등식뿐 (U·용량차이 아님) → (a) 캐시 지키려 GPU 굶김(R<1) · (b) 큰 GPU 덜 채움(:358)")

# ===========================================================================
# 10 — 증명 ① 실험 C (인과, R 모델)
# ===========================================================================
s = slide()
title(s, "증명 ①  실험 C — 인과 확정 (완료)", GREEN, "앞의 '한 식'을 통제 실험으로 증명 — 예측 U vs 실측 U  r = 0.96")
add_image(s, "expC_R1_crossing.png", 1.50, width=11.5, left=0.90)
textbox(s, 0.90, 4.84, 11.5, 0.30,
        [[("duty 축으로 R을 0.31→1.0으로 밀자 U·throughput이 예측대로 회복 (R=1에서 교차)", 11, GRAY, False)]],
        align=PP_ALIGN.CENTER)
# left panel: R model + numbers
rect(s, 0.90, 5.18, 7.15, 1.72, PANEL, line=GREEN, lw=1.3, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
textbox(s, 1.10, 5.26, 6.85, 1.6,
        [[("R 모델:  U ≈ min(R, 1),   R = k_fit × d", 13, GREEN, True),
          ("     (예측 U vs 실측 U  r = 0.959)", 10.5, GRAY, False)],
         [("tr: R≈0.31 <1 → U≈0.35 (GPU 65% 놀음) → throughput 열세", 11.5, RED, True)],
         [("인과 결정타: R을 1까지 밀면 throughput 0.068→0.144(×2.1),", 11.5, INK, True)],
         [("   hit 유지 → tr이 default를 두 축 모두 역전", 11.5, GREEN, True)]])
# right panel: caveat
rect(s, 8.25, 5.18, 4.5, 1.72, AMBRL, line=AMBER, lw=1.2, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
textbox(s, 8.45, 5.26, 4.2, 1.6,
        [[("주의 — 두 종류의 회복", 12, AMBER, True)],
         [("pause만 완화:", 11, RED, True), (" thru↑ 지만", 10.5, INK, False)],
         [("   hit 0.79→0.65 (나쁜 회복)", 10.5, RED, False)],
         [("duty↑ / 용량↑:", 11, GREEN, True), (" thru↑ +", 10.5, INK, False)],
         [("   hit 유지 (좋은 회복)", 10.5, GREEN, True)]])
notes(s,
"이제 '이게 진짜 원인이 맞다'를 실험으로 증명한 부분입니다. 실험 C, 완료된 결과입니다. 핵심은 R이라는 "
"하나의 값입니다. tr의 GPU 활용률 U는 R과 거의 같은데, R은 동시에 올라간 프로그램 수 k_fit 곱하기 워크로드 "
"duty d입니다. duty는 전체 시간 중 GPU가 실제로 도는 비율이라고 보시면 됩니다. 예측한 U와 실측 U의 상관이 "
"0.959로 거의 완벽합니다. tr은 캐시를 지키려고 pause를 많이 해서 k_fit이 낮게 유지되고, 그래서 R이 약 "
"0.31로 1보다 작습니다. R이 1보다 작으면 GPU가 그만큼 놉니다. 여기서 결정적인 실험 — 워크로드 duty를 "
"조절해서 R을 1까지 밀어봤더니, throughput이 0.068에서 0.144로 두 배 넘게 오르면서 default를 처리량과 "
"hit rate 두 축 모두에서 역전했습니다. 즉 'R이 1보다 작은 게 tr이 지는 원인'이라는 걸 인과적으로 확정한 "
"겁니다. 오른쪽 아래 주의점 하나 — pause만 완화해도 처리량은 오르지만 그때는 hit rate가 떨어집니다. 이건 "
"'나쁜 회복'이고, 진짜 해법은 duty를 올리거나 용량을 키우는 '좋은 회복'입니다.")
footer(s, "R<1이 tr 열세의 원인 — duty로 R을 1까지 밀면 hit 지키며 throughput 역전 (r=0.96)")

# ===========================================================================
# 10 — 증명 ② SWE-bench (2차 실데이터)
# ===========================================================================
s = slide()
title(s, "증명 ②  SWE-bench — 2차 실데이터 (완료)", GREEN, "duty가 다른 워크로드에서 R 모델이 승패를 미리 맞힌다")
add_image(s, "swebench_throughput.png", 1.62, width=5.55, left=0.72)
add_image(s, "swebench_hitrate.png", 1.62, width=5.55, left=6.55)
# contrast band — 세 워크로드, 한 식
rect(s, 0.90, 5.83, 11.85, 1.14, PANEL, line=GREEN, lw=1.3, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
textbox(s, 1.10, 5.90, 11.5, 1.04,
        [[("같은 2×4090 — duty(d) 하나가 R을, R이 승패를 결정 (세 워크로드 · 한 식)", 13, INK, True)],
         [("합성 d≈0.56 → R≈1.7 → 승 (+57%)    ·    ", 11.5, GREEN, True),
          ("TraceLab d≈0.18 → R<1 → 패 (−34%)    ·    ", 11.5, RED, True),
          ("SWE d≈1.0 → R≫1 → 승 (+78~84%)", 11.5, GREEN, True)],
         [("SWE는 스윕 전에 R로 예측 → 적중. 흩어진 승패가 하나의 값으로 통일.", 11, GRAY, True)]])
notes(s,
"두 번째 증명은 완전히 다른 실데이터인 SWE-bench입니다. 이것도 완료된 결과입니다. SWE-bench는 코딩 "
"에이전트라 생성이 길어서 duty d가 거의 1입니다. 아까 TraceLab은 d가 0.18이었죠. R 모델대로라면 d가 크니까 "
"R이 1보다 훨씬 커지고, 그러면 tr이 이겨야 합니다. 그리고 저는 이 예측을 스윕을 돌리기 전에 미리 적어놨고, "
"실제로 맞았습니다. 왼쪽 throughput — 부하가 높을 때 tr이 default보다 78%에서 84% 더 높습니다. 앞의 "
"TraceLab에서 34% 졌던 것과 정반대입니다. 오른쪽 hit rate는 여기서도 tr이 0.8 근처를 지키고 default는 "
"붕괴합니다. 아래 띠가 오늘 발표의 하이라이트입니다 — 똑같은 2×4090 하드웨어인데, 한 워크로드에선 tr이 "
"34% 지고 다른 워크로드에선 84% 이깁니다. 이 정반대 결과를 가르는 유일한 변수가 워크로드의 duty 하나이고, "
"그걸 R 모델이 미리 예측했습니다. 흩어져 있던 승패가 하나의 값으로 통일된 겁니다.")
footer(s, "같은 HW, duty만 다름: R 모델이 승패를 사전 예측 → 두 실데이터로 확증")

# ===========================================================================
# 11 — 향후 계획
# ===========================================================================
s = slide()
title(s, "향후 계획", BLUE, "원인을 알았으니, 이종·분산 환경의 해법으로")

plans = [
    ("① Related work 추가 조사", AMBER,
     ["Continuum — 툴 콜 중 KV cache TTL 유지로 재프리필 회피",
      "DistServe — prefill/decode 분리, 그 외 이종 서빙 기법 폭넓게"]),
    ("② 용량 비례 라우팅", GREEN,
     ["신규 배정을 절대 토큰이 아닌 용량 비례로 (router.py:358 핵심)",
      "5090 과소활용 해소 → 이종에서 default 역전 목표"]),
    ("③ 이종 → 분산으로 확장", BLUE,
     ["R 모델을 이종/분산 클러스터의 라우팅 설계 원리로 일반화",
      "duty·용량을 함께 보는 스케줄링 정책 탐색"]),
]
y = 1.75
for h, col, items in plans:
    rect(s, 0.90, y, 11.85, 1.55, PANEL, line=col, lw=1.4, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    rect(s, 0.90, y, 0.14, 1.55, col)
    textbox(s, 1.20, y+0.12, 11.3, 0.45, [[(h, 14.5, col, True)]])
    textbox(s, 1.20, y+0.60, 11.3, 0.9,
            [[("• "+items[0], 12, INK, False)], [("• "+items[1], 12, INK, False)]])
    y += 1.72
notes(s,
"마지막으로 향후 계획입니다. 원인을 알았으니 이제 해법으로 갑니다. 첫째, related work를 더 조사하겠습니다. "
"특히 Continuum이라는 연구가 툴 콜 중에 KV 캐시를 TTL로 유지해서 재프리필을 피하는데, 제 문제의식과 정확히 "
"맞닿아 있어 깊게 볼 계획입니다. DistServe 같은 프리필/디코드 분리 기법도 함께 보겠습니다. 둘째, 가장 직접적인 "
"해법으로 용량 비례 라우팅을 시도하겠습니다. 앞서 5090 과소활용의 원인이 router.py 358번 라인의 절대 토큰 "
"배정이었으니, 이걸 용량 비례로 바꾸면 이종에서 default를 역전할 수 있을 것으로 봅니다. 셋째, 이걸 이종을 "
"넘어 분산 클러스터까지 확장할 수 있는지 탐색하겠습니다. R 모델을 라우팅 설계의 일반 원리로 삼아, duty와 "
"용량을 함께 고려하는 정책을 만드는 게 목표입니다. 이상입니다. 감사합니다.")
footer(s, "Continuum/DistServe 조사 · 용량 비례 라우팅 · 이종→분산 확장")

# ===========================================================================
# 12 — 요약 한 장
# ===========================================================================
s = slide()
rect(s, 0, 0, 13.333, 7.5, WHITE)
rect(s, 0.80, 0.75, 0.14, 0.75, BLUE)
textbox(s, 1.05, 0.72, 11.5, 0.8, [[("한 장 요약", 25, INK, True)]], anchor=MSO_ANCHOR.MIDDLE)
rect(s, 1.05, 1.62, 11.5, 0.02, BLUE)

lines = [
    ("문제", RED, "에이전트 서빙 — 멀티턴 KV 축적·툴 버블·스래싱·로컬리티 vs 밸런싱, 게다가 기존 연구는 동종 GPU 가정"),
    ("접근", AMBER, "SOTA인 ThunderAgent(tr)를 baseline으로, 2×4090에서 워크로드 duty만 바꿔가며 병목 탐색"),
    ("결과", BLUE, "tr은 캐시 hit는 항상 지킴. throughput은 합성 +57%(승) → TraceLab −34%(패)로 뒤집힘"),
    ("원인", INK, "스케줄러가 '용량 부등식(스래싱 회피)'만 판단 → 캐시 지키려 GPU 굶김(R<1) / 큰 GPU 덜 채움(절대토큰)"),
    ("증명", GREEN, "U≈min(R,1), R=k_fit×d 하나로 승패 통일(r=0.96). 실험 C 인과 확정, SWE 사전예측 적중 — 둘 다 완료"),
    ("향후", BLUE, "용량 비례 라우팅 + Continuum 등 조사 + 이종→분산 확장"),
]
yy = 1.95
for tag, col, txt in lines:
    boxtext(s, 1.05, yy, 1.55, 0.72, [[(tag, 14, WHITE, True)]], col)
    textbox(s, 2.80, yy, 9.85, 0.72, [[(txt, 13, INK, False)]], anchor=MSO_ANCHOR.MIDDLE)
    yy += 0.87
notes(s,
"한 장으로 정리하면 이렇습니다. 문제는 에이전트 서빙의 네 가지 어려움에 더해 기존 연구가 동종 GPU를 "
"가정한다는 점, 접근은 SOTA인 ThunderAgent를 baseline으로 이종 환경에서 병목을 찾은 것, 결과는 tr이 "
"캐시는 항상 지키지만 처리량은 워크로드에 따라 지기도 이기기도 한다는 것입니다. 원인은 논문의 비용 모델을 "
"실제로 최적화하지 않고 근사한 데서 왔고, 그걸 R이라는 하나의 값으로 통일해서 인과까지 증명했습니다. "
"앞으로는 용량 비례 라우팅과 분산 확장으로 갑니다. 감사합니다. 질문 받겠습니다.")
footer(s, "핵심 한 줄 — 흩어진 승패를 R=k_fit×d 하나로 설명하고 인과까지 확정했다")

prs.save(OUT)
print("SAVED", OUT)
print("slides:", len(prs.slides._sldIdLst))
